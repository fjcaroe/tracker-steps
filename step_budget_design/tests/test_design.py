import base64,io
from types import SimpleNamespace
import openpyxl
from lxml import etree
from odoo import Command
from odoo.exceptions import UserError,ValidationError
from odoo.tests.common import TransactionCase,new_test_user
from ..models.excel import ExcelValues
from ..models.template import season_months


class BudgetDesignCase(TransactionCase):
    def sample(self):
        env=self.env(context=dict(self.env.context,tracking_disable=True,mail_create_nosubscribe=True,no_reset_password=True))
        season=env['step.temporada'].create({'name':'QA T77 real calendar','start_date':'2026-09-01','end_date':'2027-08-31'})
        version=env['step.management.budget.version'].create({'code':'QA-'+str(season.id),'name':'QA T77 version'})
        species=env['step.especie'].create({'name':'QA T77 species','type_especie':'frutal','group_especie':'baya'})
        plan=env['account.analytic.plan'].create({'name':'QA T77 centers'})
        activity_plan=env['account.analytic.plan'].create({'name':'Actividad'})
        activity=env['account.analytic.account'].create({'name':'QA T77 pruning','plan_id':activity_plan.id})
        center=env['account.analytic.account'].create({'name':'QA T77 4ha','plan_id':plan.id,'hectares':4,'company_id':env.company.id})
        group=env['step.management.budget.group'].create({'name':'QA T77 own labor','code':'QA-'+str(season.id),'origin_id':env.ref('step_budget_design.origin_labor').id})
        uom=env.ref('uom.product_uom_day')
        product=env['product.product'].create({'name':'QA T77 pruning','uom_id':uom.id,'uom_po_id':uom.id,'actividad_id':activity.id})
        template=env['step.management.budget.template'].create({'name':'QA T77 template','season_id':season.id,'budget_version_id':version.id,'species_id':species.id,
            'base_hectares':1,'base_plants':3000,'base_kg':12000,'line_ids':[Command.create({'indicator':product.name,'origin_id':env.ref('step_budget_design.origin_labor').id,
                'group_id':group.id,'activity_id':activity.id,'product_id':product.id,'uom_id':uom.id,'calculation_uom_id':env.ref('step_budget_design.calculation_plants').id,
                'yield_quantity':200,'unit_price':150,'month_ids':[Command.create({'period_date':'2026-09-01','quantity':15})]})]})
        budget=env['step.management.operational.budget'].create({'description':'QA T77 budget','template_id':template.id,
            'allocation_ids':[Command.create({'center_id':center.id,'hectares':4})]})
        return template,budget

    def test_pruning_unitary_quantity_and_hectare_application(self):
        template,budget=self.sample();line=template.line_ids
        self.assertEqual(line.calculation_quantity,15);self.assertEqual(line.calculation_amount,450000);self.assertEqual(line.calculation_unit_price,30000)
        self.assertIn('2026-09',template.quantities_matrix);self.assertIn('2027-08',template.values_matrix)
        budget.action_generate_lines();self.assertEqual(budget.line_ids.quantity,60);self.assertEqual(budget.total_amount,1800000)
        self.assertEqual(budget.line_ids.month_ids.period_date,template.line_ids.month_ids.period_date)
        self.assertEqual(budget.line_ids.month_ids.conversion_date.year,2026)
        self.assertTrue(budget.line_ids.distribution_complete)

    def test_harvest_hectares_days_units_and_income(self):
        template,budget=self.sample();line=template.line_ids
        line.write({'calculation_uom_id':self.env.ref('step_budget_design.calculation_kg').id,'yield_quantity':50,'unit_price':600})
        self.assertEqual(line.calculation_quantity,240);self.assertEqual(line.calculation_amount,7200000);self.assertEqual(line.calculation_unit_price,30000)
        line.write({'calculation_uom_id':self.env.ref('step_budget_design.calculation_hectares').id,'yield_quantity':2.5,'unit_price':6500})
        self.assertEqual(line.calculation_quantity,2.5);self.assertEqual(line.calculation_amount,16250)
        line.write({'calculation_uom_id':self.env.ref('step_budget_design.calculation_days').id,'base_quantity':10,'unit_price':35000})
        self.assertEqual(line.calculation_amount,350000)
        group=self.env['step.management.budget.group'].create({'code':'QA-INCOME-'+str(template.id),'name':'QA Income','flow_type':'income'})
        line.write({'group_id':group.id,'calculation_uom_id':self.env.ref('step_budget_design.calculation_kg').id,'unit_price':2000,'yield_quantity':1})
        self.assertEqual(line.calculation_quantity,12000);self.assertEqual(line.calculation_unit_price,2000)

    def test_distribution_validation_and_template_duplicate(self):
        template,budget=self.sample();template.line_ids.month_ids.quantity=14
        with self.assertRaises(ValidationError):template.action_activate()
        template.line_ids.month_ids.quantity=15;template.action_activate()
        with self.assertRaises(UserError):template.line_ids.unit_price=1
        copy=self.env[template._name].browse(template.action_duplicate_template()['res_id'])
        self.assertEqual(copy.state,'draft');self.assertEqual(copy.line_ids.month_ids.quantity,15)
        copy.line_ids.unit_price=160
        self.assertEqual(template.line_ids.unit_price,150)

    def test_season_calendar_months_and_monthly_dollar(self):
        template,budget=self.sample();months=season_months(template.season_id)
        self.assertEqual(str(months[0]),'2026-09-01');self.assertEqual(str(months[-1]),'2027-08-01')
        template.line_ids.month_ids.write({'period_date':'2027-01-01'})
        rates=self.env['step.management.exchange.rate']
        rate=rates.search([('company_id','=',self.env.company.id),('currency_id','=',self.env.ref('base.USD').id),('year','=',2027),('month','=','01')])
        if rate:rate.company_value_per_unit=900
        else:rates.create({'company_id':self.env.company.id,'currency_id':self.env.ref('base.USD').id,'year':2027,'month':'01','company_value_per_unit':900})
        row=template.line_ids.month_ids
        expected=450000 if template.currency_id==self.env.ref('base.USD') else 500
        self.assertAlmostEqual(row.amount_usd,expected)
        budget.action_generate_lines();self.assertEqual(str(budget.line_ids.month_ids.conversion_date),'2027-01-01')

    def test_approval_revision_report_preserves_parameters(self):
        template,budget=self.sample();budget.action_generate_lines();budget.action_approve()
        revision=self.env[budget._name].browse(budget.action_new_revision()['res_id'])
        self.assertEqual(revision.season_id,budget.season_id);self.assertEqual(revision.line_ids.activity_id,budget.line_ids.activity_id)
        self.assertEqual(revision.line_ids.month_ids.period_date,budget.line_ids.month_ids.period_date)
        with self.assertRaises(UserError):budget.line_ids.month_ids.period_date='2026-10-01'
        self.env.flush_all()
        summary=self.env['step.management.budget.summary.report'].search([('budget_id','=',budget.id)])
        self.assertEqual(summary.hectares,4);self.assertEqual(summary.expected_kg,48000);self.assertEqual(summary.cost,1800000)
        self.assertEqual(summary.cost_per_ha,450000)
        self.assertEqual(summary.jornadas_per_ha,15)
        data=self.env['step.management.budget.data.report'].search([('budget_id','=',budget.id)])
        self.assertEqual(data.quantity,60);self.assertEqual(data.jornadas,60)

    def test_general_direct_amounts_without_hectare_template(self):
        template,budget=self.sample();budget.write({'budget_type':'general','template_id':False,'season_id':template.season_id.id,
            'line_ids':[Command.create({'center_id':budget.allocation_ids.center_id.id,'group_id':template.line_ids.group_id.id,'indicator':'QA direct','calculation_mode':'direct','direct_amount':15000,
                'month_ids':[Command.create({'month':'sep','period_date':'2026-09-01','direct_amount':15000})]})]})
        budget.action_generate_lines();self.assertEqual(budget.state,'calculated');self.assertEqual(budget.total_amount,15000)
        budget.action_approve();self.assertEqual(budget.state,'approved')

    def test_machinery_components_rate_and_zero_hour_guard(self):
        template,budget=self.sample();env=template.env
        brand=env['fleet.vehicle.model.brand'].create({'name':'QA T77 brand'})
        model=env['fleet.vehicle.model'].create({'name':'QA T77 tractor','brand_id':brand.id})
        vehicle=env['fleet.vehicle'].create({'model_id':model.id,'license_plate':'QA-T77','company_id':env.company.id})
        service=env['type.service.machinery'].search([],limit=1)
        if not service:service=env['type.service.machinery'].create({'name':'QA Fuel','cod':'QA77'})
        source=env['step.management.machinery.budget'].create({'name':'QA machinery','season_id':template.season_id.id,'budget_version_id':template.budget_version_id.id,'vehicle_id':vehicle.id,
            'month_ids':[Command.create({'period_date':p,'hours':10}) for p in season_months(template.season_id)],
            'component_ids':[Command.create({'period_date':'2026-09-01','service_type_id':service.id,'consumption_per_hour':4,'unit_price':1000}),
                Command.create({'period_date':'2026-09-01','service_type_id':service.id,'calculation_mode':'direct','direct_amount':20000})]})
        self.assertEqual(source.total_hours,120);self.assertEqual(source.total_cost,60000);self.assertEqual(source.hourly_rate,500)
        source.action_approve();template.line_ids.machinery_budget_id=source;template.line_ids.action_use_machinery_budget()
        self.assertEqual(template.line_ids.unit_price,500)
        with self.assertRaises(UserError):source.component_ids[0].unit_price=1

    def test_safe_formula_reader_and_formula_injection_rejected(self):
        book=openpyxl.Workbook();sheet=book.active;sheet.title='datos';sheet['B11']=3000;sheet['B12']='=+$B$11*4';sheet['F27']=200;sheet['I27']='=+$B$11/F27'
        reader=ExcelValues(book);self.assertEqual(reader.value('datos','B12'),12000);self.assertEqual(reader.value('datos','I27'),15)
        sheet['A1']='=__import__("os").system("whoami")'
        with self.assertRaises(ValidationError):reader.value('datos','A1')

    def test_normalized_excel_direct_amounts_keep_both_months(self):
        template,budget=self.sample();line=template.line_ids
        book=openpyxl.Workbook();sheet=book.active
        sheet.append(['Version Ppto','Temporada','Ano','Mes','Fundo','Especie','Variedad','Tipo CCosto','Centro de costos','Origen','Grupo Presupuesto','Actividad','Producto-labor','UdM','Cantidad','Valor Ppto$','TC Ppto','Valor Ppto US$'])
        for month,amount in ((9,15000),(10,20000)):
            sheet.append(['QA',template.season_id.name,2026,month,None,None,None,'Administrativo',budget.allocation_ids.center_id.name,'Mano de obra',line.group_id.name,line.activity_id.name,line.product_id.name,line.uom_id.name,None,amount,900,amount/900])
        raw=io.BytesIO();book.save(raw)
        source=self.env['step.management.budget.import'].create({'file':base64.b64encode(raw.getvalue()),'filename':'qa-general.xlsx',
            'season_master_id':template.season_id.id,'version_master_id':template.budget_version_id.id,'target_budget_type':'general'})
        source.action_validate();self.assertEqual(source.error_count,0)
        imported=self.env['step.management.operational.budget'].browse(source.action_import()['res_id'])
        self.assertEqual(imported.budget_type,'general');self.assertEqual(imported.total_amount,35000)
        self.assertEqual(imported.line_ids.calculation_mode,'direct')
        self.assertEqual(sum(imported.line_ids.month_ids.mapped('amount')),35000)
        self.assertEqual(set(str(d) for d in imported.line_ids.month_ids.mapped('period_date')),{'2026-09-01','2026-10-01'})
        with self.assertRaises(UserError):source.action_import()

    def test_import_equivalent_agricultural_template(self):
        template,budget=self.sample();line=template.line_ids
        book=openpyxl.Workbook();s=book.active;s.title='datos';s['B10']=1;s['B11']=3000;s['B12']='=B11*4'
        for cell,value in {'A22':'Mano de obra','B22':line.group_id.name,'C22':line.activity_id.name,'D22':line.product_id.name,'E22':'Plantas','F22':200,'G22':150,'H22':line.uom_id.name,'I22':'=B11/F22'}.items():s[cell]=value
        d=book.create_sheet('distribucion cantidad');d['B6']=line.group_id.name;d['C6']=line.activity_id.name;d['D6']=line.product_id.name;d['G6']=15
        raw=io.BytesIO();book.save(raw)
        wizard=self.env['step.management.template.import.wizard'].create({'name':'QA import','season_id':template.season_id.id,'budget_version_id':template.budget_version_id.id,'species_id':template.species_id.id,'file':base64.b64encode(raw.getvalue()),'filename':'qa.xlsx'})
        imported=self.env[template._name].browse(wizard.action_import()['res_id'])
        self.assertEqual(imported.line_ids.calculation_quantity,15);self.assertEqual(imported.line_ids.calculation_unit_price,30000)
        self.assertEqual(str(imported.line_ids.month_ids[0].period_date),'2026-09-01')

    def test_roles_effective_forms_and_master_reuse(self):
        env=self.env(context=dict(self.env.context,no_reset_password=True))
        user=new_test_user(env,login='qa-t77-planner',groups='step_management_costs.group_management_manager')
        visible=env['ir.ui.menu'].with_user(user)._visible_menu_ids()
        for name in ('template_import_menu','machinery_menu','data_menu','summary_menu','days_menu','center_menu','labor_menu'):
            self.assertIn(env.ref('step_budget_design.'+name).id,visible)
        arch=etree.fromstring(env['step.management.budget.template'].with_user(user).get_view(view_id=env.ref('step_management_costs.view_budget_template_form').id,view_type='form')['arch'].encode())
        for name in ('season_id','calculation_uom_id','calculation_unit_price','period_date'):self.assertTrue(arch.xpath('.//field[@name="%s"]'%name))
        self.assertEqual(env['step.management.budget.group']._fields['phenology_id'].comodel_name,'step.agriculture.phenological.stage')
