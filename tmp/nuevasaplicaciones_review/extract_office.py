from __future__ import annotations

import sys
from pathlib import Path


def read_docx(path: Path) -> None:
    from docx import Document

    doc = Document(path)
    print(f"FILE: {path}")
    print("PARAGRAPHS:")
    for idx, paragraph in enumerate(doc.paragraphs, 1):
        text = paragraph.text.strip()
        if text:
            print(f"P{idx}: {text}")
    for section_idx, section in enumerate(doc.sections, 1):
        for kind, paragraphs in (
            ("HEADER", section.header.paragraphs),
            ("FOOTER", section.footer.paragraphs),
        ):
            for idx, paragraph in enumerate(paragraphs, 1):
                text = paragraph.text.strip()
                if text:
                    print(f"{kind}{section_idx}.{idx}: {text}")
    for table_idx, table in enumerate(doc.tables, 1):
        print(f"TABLE {table_idx} ({len(table.rows)}x{len(table.columns)}):")
        for row_idx, row in enumerate(table.rows, 1):
            values = [cell.text.replace("\n", " | ").strip() for cell in row.cells]
            print(f"R{row_idx}: " + " || ".join(values))


def read_xlsx(path: Path) -> None:
    import openpyxl

    book = openpyxl.load_workbook(path, data_only=False, read_only=False)
    print(f"FILE: {path}")
    for sheet in book.worksheets:
        print(f"SHEET: {sheet.title} DIMENSION={sheet.calculate_dimension()}")
        for row in sheet.iter_rows():
            populated = []
            for cell in row:
                if cell.value is not None:
                    populated.append(f"{cell.coordinate}={cell.value!r}")
            if populated:
                print(" | ".join(populated))


def main() -> None:
    for arg in sys.argv[1:]:
        path = Path(arg)
        if path.suffix.lower() == ".docx":
            read_docx(path)
        elif path.suffix.lower() == ".xlsx":
            read_xlsx(path)
        else:
            raise SystemExit(f"Unsupported: {path}")
        print("\n" + "=" * 100 + "\n")


if __name__ == "__main__":
    main()
