"""Corte V2 B — cargas históricas oficiales (presupuesto y real) y
plantillas versionadas (`PROMPT_CLAUDE_SONNET5_RESPUESTAS_V2_2026-09-06.md`).

No usa los archivos reales del cliente (16.045 / 18.638 filas, datos del
cliente) como fixture: construye muestras representativas pequeñas con
exactamente la misma cabecera de las plantillas oficiales
(`data/historical_template_versions.xml`), reproduciendo cada escenario
exigido por la puerta de salida.
"""

import base64
import io

import openpyxl
from psycopg2 import IntegrityError

from odoo.exceptions import UserError, ValidationError
from odoo.tests import tagged
from odoo.tools import mute_logger

from .test_management_costs import ManagementCostsCommon

BUDGET_HEADERS = [
    "Versión Ppto", "Temporada", "Año", "Mes", "Fundo", "Especie", "Variedad",
    "Tipo CCosto", "Centro de costos", "Origen", "Grupo Presupuesto",
    "Actividad", "Producto-labor", "UdM", "Cantidad", "Valor Ppto$", "TC Ppto",
    "Valor Ppto US$",
]
_UNSET = object()

ACTUAL_HEADERS = [
    "Tipo registro", "Temporada", "Año", "Mes", "Fundo", "Especie", "Variedad",
    "Tipo CCosto", "Centro de costos", "Origen", "Grupo Presupuesto",
    "Actividad", "Producto-labor", "UdM", "Cantidad", "Valor Real $", "TC Real",
    "Valor Real US$",
]


def _xlsx(headers, rows, formula_rates=None):
    """`rows`: lista de listas (18 valores). `formula_rates`: {row_index(1-based
    dentro de rows) : formula_string} para forzar una fórmula en la columna TC."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(headers)
    for row in rows:
        ws.append(row)
    for idx, formula in (formula_rates or {}).items():
        ws.cell(row=idx + 1, column=17).value = formula  # +1 por la cabecera
    bio = io.BytesIO()
    wb.save(bio)
    return base64.b64encode(bio.getvalue())


@tagged("post_install", "-at_install")
class TestV2BHistoricalImport(ManagementCostsCommon):
    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.center_a.write({"cost_type": "crop"})
        cls.center_operational = cls.env["step.management.cost.center"].create({
            "code": "OPV2B", "name": "Operacional V2B", "company_id": cls.company_a.id,
            "cost_type": "operational", "hectares": 0.0,
        })

    def _budget_row(self, version="1", season="2425", year=2024, month=11,
                     farm="Fundo A", species="Cerezo", variety="Bing",
                     cost_type="Frutales", center=None, origin="Mano de Obra",
                     group=_UNSET, activity="Poda", product="Labor poda",
                     uom="Jornada", quantity=10, amount=100000.0, rate=900.0,
                     amount_usd=None):
        center = center or self.center_a
        group = self.group_a if group is _UNSET else group
        if amount_usd is None:
            amount_usd = round(amount / rate, 4)
        return [
            version, season, year, month, farm, species, variety, cost_type,
            center.code, origin, group.code if group else "", activity, product,
            uom, quantity, amount, rate, amount_usd,
        ]

    def _actual_row(self, record_type="Externo", **kw):
        row = self._budget_row(**kw)
        row[0] = record_type
        return row

    def _make_budget_batch(self, b64, **kw):
        vals = {
            "company_id": self.company_a.id, "dataset_kind": "budget",
            "file": b64, "filename": "presupuesto.xlsx",
        }
        vals.update(kw)
        return self.env["step.management.historical.import.batch"].create(vals)

    def _make_actual_batch(self, b64, **kw):
        vals = {
            "company_id": self.company_a.id, "dataset_kind": "actual",
            "file": b64, "filename": "real.xlsx",
        }
        vals.update(kw)
        return self.env["step.management.historical.import.batch"].create(vals)

    # ------------------------------------------------------------------
    # 18 columnas / plantilla vigente
    # ------------------------------------------------------------------
    def test_budget_18_columns_happy_path(self):
        b64 = _xlsx(BUDGET_HEADERS, [self._budget_row()])
        batch = self._make_budget_batch(b64)
        batch.action_validate()
        self.assertEqual(batch.error_count, 0, batch.line_ids.mapped("error"))
        self.assertTrue(batch.template_version_id)
        batch.action_import()
        self.assertEqual(batch.state, "entered")
        fact = batch.cost_ids
        self.assertEqual(len(fact), 1)
        self.assertEqual(fact.dataset_kind, "budget")
        self.assertEqual(fact.season, "2024/2025")
        self.assertAlmostEqual(fact.budget_amount, 100000.0)
        self.assertEqual(fact.center_id, self.center_a)
        self.assertEqual(fact.origin, "external")

    def test_unknown_template_version_rejected(self):
        headers = BUDGET_HEADERS + ["Columna extra"]
        rows = [self._budget_row() + ["x"]]
        b64 = _xlsx(headers, rows)
        batch = self._make_budget_batch(b64)
        with self.assertRaises(UserError):
            batch.action_validate()

    # ------------------------------------------------------------------
    # Fórmula de TC: segura vs. no permitida
    # ------------------------------------------------------------------
    def test_actual_safe_tc_formula_computed_in_server(self):
        row = self._actual_row(amount=100000.0, amount_usd=111.11, rate=0.0)
        b64 = _xlsx(ACTUAL_HEADERS, [row], formula_rates={1: "=P2/R2"})
        batch = self._make_actual_batch(b64)
        batch.action_validate()
        self.assertEqual(batch.error_count, 0, batch.line_ids.mapped("error"))
        line = batch.line_ids
        self.assertAlmostEqual(line.source_exchange_rate, 100000.0 / 111.11, places=2)

    def test_actual_unsafe_formula_rejected(self):
        row = self._actual_row(amount=100000.0, amount_usd=111.11, rate=0.0)
        b64 = _xlsx(ACTUAL_HEADERS, [row], formula_rates={1: "=P2*2"})
        batch = self._make_actual_batch(b64)
        batch.action_validate()
        self.assertEqual(batch.error_count, 1)
        self.assertIn("fórmula", batch.line_ids.error.lower())

    def test_formula_in_other_column_rejected(self):
        row = self._budget_row()
        b64 = _xlsx(BUDGET_HEADERS, [row])
        # inyecta una fórmula en "Cantidad" (columna O = 15) directamente en el xml
        wb = openpyxl.load_workbook(io.BytesIO(base64.b64decode(b64)))
        wb.active.cell(row=2, column=15).value = "=1+1"
        bio = io.BytesIO()
        wb.save(bio)
        b64_bad = base64.b64encode(bio.getvalue())
        batch = self._make_budget_batch(b64_bad)
        batch.action_validate()
        self.assertEqual(batch.error_count, 1)
        self.assertIn("fórmula", batch.line_ids.error.lower())

    # ------------------------------------------------------------------
    # Especie obligatoria según tipo de centro
    # ------------------------------------------------------------------
    def test_species_required_for_crop_center(self):
        row = self._budget_row(species="", center=self.center_a)
        b64 = _xlsx(BUDGET_HEADERS, [row])
        batch = self._make_budget_batch(b64)
        batch.action_validate()
        self.assertEqual(batch.error_count, 1)
        self.assertIn("especie", batch.line_ids.error.lower())

    def test_species_optional_for_operational_center(self):
        row = self._budget_row(
            species="", variety="", cost_type="Operacional", center=self.center_operational,
        )
        b64 = _xlsx(BUDGET_HEADERS, [row])
        batch = self._make_budget_batch(b64)
        batch.action_validate()
        self.assertEqual(batch.error_count, 0, batch.line_ids.mapped("error"))

    # ------------------------------------------------------------------
    # Maestro inexistente y ambiguo
    # ------------------------------------------------------------------
    def test_unknown_center_is_row_error(self):
        row = self._budget_row()
        row[8] = "NoExiste123"
        b64 = _xlsx(BUDGET_HEADERS, [row])
        batch = self._make_budget_batch(b64)
        batch.action_validate()
        self.assertEqual(batch.error_count, 1)
        self.assertIn("no encontrado", batch.line_ids.error.lower())

    def test_ambiguous_center_is_row_error(self):
        self.env["step.management.cost.center"].create({
            "code": "AMBV2B1", "name": "Ambiguo V2B", "company_id": self.company_a.id,
        })
        self.env["step.management.cost.center"].create({
            "code": "AMBV2B2", "name": "ambiguo v2b", "company_id": self.company_a.id,
        })
        row = self._budget_row()
        row[8] = "AMBIGUO V2B"  # no calza exacto con ninguno; normalizado calza con ambos
        b64 = _xlsx(BUDGET_HEADERS, [row])
        batch = self._make_budget_batch(b64)
        batch.action_validate()
        self.assertEqual(batch.error_count, 1)
        self.assertIn("ambigu", batch.line_ids.error.lower())

    # ------------------------------------------------------------------
    # Duplicado exacto dentro del lote
    # ------------------------------------------------------------------
    def test_exact_duplicate_shown_and_included_by_default(self):
        row = self._budget_row()
        b64 = _xlsx(BUDGET_HEADERS, [row, list(row)])
        batch = self._make_budget_batch(b64)
        batch.action_validate()
        self.assertEqual(batch.duplicate_count, 1)
        batch.action_import()
        self.assertEqual(len(batch.cost_ids), 2)  # por omisión, se incluyen

    def test_exact_duplicate_can_be_excluded(self):
        row = self._budget_row()
        b64 = _xlsx(BUDGET_HEADERS, [row, list(row)])
        batch = self._make_budget_batch(b64, exclude_exact_duplicates=True)
        batch.action_validate()
        batch.action_import()
        self.assertEqual(len(batch.cost_ids), 1)
        # la fila duplicada sigue visible en el detalle (auditoría)
        self.assertEqual(batch.line_count, 2)

    # ------------------------------------------------------------------
    # Mismo archivo concurrente / dos compañías
    # ------------------------------------------------------------------
    def test_same_file_concurrent_blocked(self):
        b64 = _xlsx(BUDGET_HEADERS, [self._budget_row()])
        batch1 = self._make_budget_batch(b64)
        batch1.action_validate()
        batch1.action_import()
        batch2 = self._make_budget_batch(b64)
        batch2.action_validate()
        with self.assertRaises(UserError):
            batch2.action_import()

    def test_two_companies_isolated(self):
        row_a = self._budget_row()
        b64_a = _xlsx(BUDGET_HEADERS, [row_a])
        batch_a = self._make_budget_batch(b64_a)
        batch_a.action_validate()
        batch_a.action_import()

        row_b = self._budget_row(center=self.center_b, group=None)
        b64_b = _xlsx(BUDGET_HEADERS, [row_b])
        batch_b = self._make_budget_batch(b64_b, company_id=self.company_b.id)
        batch_b.action_validate()
        self.assertEqual(batch_b.error_count, 0, batch_b.line_ids.mapped("error"))
        batch_b.action_import()
        self.assertEqual(batch_b.cost_ids.company_id, self.company_b)
        self.assertNotEqual(batch_a.company_id, batch_b.company_id)

    # ------------------------------------------------------------------
    # Conciliación monetaria
    # ------------------------------------------------------------------
    def test_amount_reconciliation_tolerance(self):
        row_ok = self._budget_row(amount=90000.0, rate=900.0, amount_usd=100.0)
        row_bad = self._budget_row(amount=90000.0, rate=900.0, amount_usd=500.0)
        b64 = _xlsx(BUDGET_HEADERS, [row_ok, row_bad])
        batch = self._make_budget_batch(b64)
        batch.action_validate()
        self.assertEqual(batch.error_count, 1)
        self.assertIn("concilia", batch.line_ids.filtered(lambda l: l.state == "error").error)

    # ------------------------------------------------------------------
    # Aprobación / inmutabilidad / reversa
    # ------------------------------------------------------------------
    def test_approval_locks_facts_and_blocks_edits(self):
        b64 = _xlsx(ACTUAL_HEADERS, [self._actual_row()])
        batch = self._make_actual_batch(b64)
        batch.action_validate()
        batch.action_import()
        self.assertEqual(batch.cost_ids.origin, "external")
        with self.assertRaises(UserError):
            batch.with_user(self.user_operator).action_approve()
        batch.with_user(self.user_approver).action_approve()
        self.assertEqual(batch.state, "approved")
        self.assertTrue(batch.cost_ids.locked)
        with self.assertRaises(UserError):
            batch.cost_ids.with_user(self.user_operator).write({"quantity": 1.0})

    def test_reversal_creates_new_batch_without_editing_approved(self):
        b64 = _xlsx(ACTUAL_HEADERS, [self._actual_row(amount=50000.0)])
        batch = self._make_actual_batch(b64)
        batch.action_validate()
        batch.action_import()
        batch.with_user(self.user_approver).action_approve()
        original_amount = batch.cost_ids.actual_amount

        reversal = batch.with_user(self.user_approver).action_create_reversal()
        reversal_batch = self.env["step.management.historical.import.batch"].browse(
            reversal["res_id"]
        )
        self.assertEqual(reversal_batch.reversal_of_id, batch)
        self.assertAlmostEqual(reversal_batch.cost_ids.actual_amount, -original_amount)
        # el lote original no se tocó
        self.assertAlmostEqual(batch.cost_ids.actual_amount, original_amount)
        self.assertTrue(batch.cost_ids.locked)

    def test_block_excludes_from_comparatives(self):
        b64 = _xlsx(ACTUAL_HEADERS, [self._actual_row()])
        batch = self._make_actual_batch(b64)
        batch.action_validate()
        batch.action_import()
        batch.with_user(self.user_approver).action_block()
        self.assertEqual(batch.state, "blocked")
        self.assertEqual(batch.cost_ids.origin, "unreviewed")
        self.assertTrue(batch.cost_ids.locked)
