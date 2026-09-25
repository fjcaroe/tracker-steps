import base64
import io

import openpyxl

from odoo.exceptions import AccessError, UserError
from odoo.tests import tagged

from .test_management_costs import ManagementCostsCommon

HEADERS = [
    "Centro de costo", "Hectáreas", "Plantas", "Rendimiento (UE)", "Kilos",
    "Temporada", "Fundo", "Especie", "Variedad",
]


def _xlsx(rows, header_at=1, trailing_junk=False):
    wb = openpyxl.Workbook()
    ws = wb.active
    for _ in range(header_at - 1):
        ws.append(["Título del reporte"] + [""] * 8)
    ws.append(HEADERS)
    for r in rows:
        ws.append(r)
    if trailing_junk:
        ws.append([""] * 9)
        ws.append([""] * 9)
        for _ in range(20):
            ws.append(["basura residual"] + [""] * 8)
    bio = io.BytesIO()
    wb.save(bio)
    return base64.b64encode(bio.getvalue())


@tagged("post_install", "-at_install")
class TestFase3EstimationImport(ManagementCostsCommon):
    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.unit = cls.env["step.management.estimation.unit"].create({
            "name": "Bin", "code": "BIN", "kg_factor": 400.0,
            "company_id": cls.company_a.id,
        })
        cls.version = cls.env["step.management.estimation.version"].create({
            "code": "V1", "name": "Estimación 2026/2027", "season": "2026/2027",
            "company_id": cls.company_a.id,
        })
        cls.category = cls.env["step.management.fruit.category"].create({
            "name": "Exportación", "code": "EXP", "company_id": cls.company_a.id,
        })
        cls.fc1 = cls.env["step.management.fruit.class"].create({
            "name": "Export", "code": "E", "category_id": cls.category.id,
            "company_id": cls.company_a.id,
        })
        cls.cg1 = cls.env["step.management.caliber.group"].create({
            "name": "Jumbo", "code": "J", "company_id": cls.company_a.id,
        })
        cls.week_curve = cls._mk_curve("week", "WK", [
            {"week_number": 45, "percentage": 50.0},
            {"week_number": 46, "percentage": 50.0},
        ])
        cls.caliber_curve = cls._mk_curve("caliber", "CB", [
            {"caliber_group_id": cls.cg1.id, "percentage": 100.0},
        ])
        cls.class_curve = cls._mk_curve("class", "CL", [
            {"fruit_class_id": cls.fc1.id, "percentage": 100.0},
        ])
        cls.center_a.write({"plants": 1000.0, "farm": "Fundo A", "plot": "C1",
                            "species": "Cerezo", "variety": "Santina"})
        cls.center_a2.write({"plants": 500.0})
        cls.center_a_no_aa.write({"plants": 0.0})

    @classmethod
    def _mk_curve(cls, curve_type, code, lines):
        curve = cls.env["step.management.estimation.curve"].create({
            "name": "Curva %s" % code, "code": code, "curve_type": curve_type,
            "company_id": cls.company_a.id,
            "line_ids": [(0, 0, v) for v in lines],
        })
        curve.action_validate()
        return curve

    def _row(self, center="CA01", hectares="", plants="", yield_ue="1", kg="",
             season="2026/2027", farm="", species="", variety=""):
        return [center, hectares, plants, yield_ue, kg, season, farm, species, variety]

    def _make_import(self, b64, method="plants", filename="est.xlsx", **kw):
        vals = {
            "company_id": self.company_a.id,
            "file": b64, "filename": filename,
            "version_id": self.version.id,
            "unit_id": self.unit.id,
            "method": method,
            "default_yield_ue": 1.0,
        }
        vals.update(kw)
        return self.env["step.management.estimation.import"].create(vals)

    # ------------------------------------------------------------------
    def test_import_plants_happy_path(self):
        imp = self._make_import(_xlsx([
            self._row(center="CA01", plants="1000", yield_ue="1.2"),
            self._row(center="CA02", plants="500", yield_ue="1"),
            self._row(center="CA03", yield_ue="0.8"),  # plantas desde el centro (0)
        ]))
        imp.action_validate()
        self.assertEqual(imp.error_count, 0, imp.line_ids.mapped("error"))
        self.assertEqual(imp.state, "validated")
        action = imp.action_import()
        self.assertEqual(imp.state, "imported")
        est = self.env["step.management.estimation"].browse(action["res_id"])
        self.assertEqual(est.state, "draft")
        self.assertEqual(est.import_id, imp)
        self.assertEqual(est.method, "plants")
        self.assertEqual(len(est.line_ids), 3)
        ca01 = est.line_ids.filtered(lambda l: l.center_id == self.center_a)
        self.assertAlmostEqual(ca01.plants, 1000.0)
        self.assertAlmostEqual(ca01.yield_ue, 1.2)
        self.assertAlmostEqual(ca01.total_ue, 1200.0)
        self.assertAlmostEqual(ca01.total_kg, 480000.0)  # 1200 * 400, factor una vez
        ca03 = est.line_ids.filtered(lambda l: l.center_id == self.center_a_no_aa)
        self.assertAlmostEqual(ca03.plants, 0.0)  # heredado del centro

    def test_import_hectares_overrides_center(self):
        imp = self._make_import(_xlsx([
            self._row(center="CA01", hectares="12", yield_ue="2"),
        ]), method="hectares")
        imp.action_validate()
        imp.action_import()
        line = imp.estimation_id.line_ids
        self.assertAlmostEqual(line.hectares, 12.0)  # del Excel, no las 10 del centro
        self.assertAlmostEqual(line.total_ue, 24.0)
        self.assertAlmostEqual(line.total_kg, 9600.0)

    def test_import_kilos_requires_kg_per_row(self):
        imp = self._make_import(_xlsx([
            self._row(center="CA01", kg="5000"),
            self._row(center="CA02"),  # sin kilos -> error
        ]), method="kilos")
        imp.action_validate()
        self.assertEqual(imp.error_count, 1)
        self.assertIn("Kilos", imp.line_ids.filtered(lambda l: l.state == "error").error)
        imp.import_valid_only = True
        imp.action_import()
        line = imp.estimation_id.line_ids
        self.assertEqual(len(line), 1)
        self.assertAlmostEqual(line.total_kg_input, 5000.0)
        self.assertAlmostEqual(line.total_kg, 5000.0)

    def test_import_all_or_nothing(self):
        imp = self._make_import(_xlsx([
            self._row(center="CA01", plants="100"),
            self._row(center="NO EXISTE", plants="100"),
        ]))
        imp.action_validate()
        self.assertEqual(imp.error_count, 1)
        before = self.env["step.management.estimation"].search_count([])
        with self.assertRaises(UserError):
            imp.action_import()
        self.assertEqual(
            self.env["step.management.estimation"].search_count([]), before
        )
        imp.import_valid_only = True
        imp.action_import()
        self.assertEqual(imp.state, "imported")
        self.assertEqual(len(imp.estimation_id.line_ids), 1)

    def test_import_idempotent_same_file(self):
        b64 = _xlsx([self._row(center="CA01", plants="100")])
        imp1 = self._make_import(b64)
        imp1.action_validate()
        imp1.action_import()
        imp2 = self._make_import(b64)
        imp2.action_validate()
        with self.assertRaises(UserError):
            imp2.action_import()

    def test_import_rejects_formula(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(HEADERS)
        ws.append(self._row(center="CA01", plants="100"))
        ws["C2"] = "=1+1"  # columna Plantas
        bio = io.BytesIO()
        wb.save(bio)
        imp = self._make_import(base64.b64encode(bio.getvalue()))
        imp.action_validate()
        self.assertEqual(imp.error_count, 1)
        self.assertIn("fórmula", (imp.line_ids.error or "").lower())

    def test_import_stops_at_residual_rows(self):
        imp = self._make_import(_xlsx(
            [self._row(center="CA01", plants="1"), self._row(center="CA02", plants="2")],
            trailing_junk=True,
        ))
        imp.action_validate()
        self.assertEqual(imp.line_count, 2)

    def test_import_duplicate_center_is_error(self):
        imp = self._make_import(_xlsx([
            self._row(center="CA01", plants="100"),
            self._row(center="CA01", plants="200"),
        ]))
        imp.action_validate()
        self.assertEqual(imp.error_count, 1)
        self.assertIn("repetido", (imp.line_ids.filtered(lambda l: l.state == "error").error or ""))

    def test_import_header_not_first_row(self):
        imp = self._make_import(_xlsx([self._row(center="CA01", plants="1")], header_at=4))
        imp.action_validate()
        self.assertEqual(imp.error_count, 0)
        self.assertEqual(imp.line_count, 1)

    def test_xlsm_is_rejected(self):
        imp = self._make_import(
            _xlsx([self._row(center="CA01", plants="1")]), filename="est.xlsm",
        )
        with self.assertRaises(UserError):
            imp.action_validate()

    def test_operator_cannot_tamper_with_validated_staging(self):
        imp = self._make_import(_xlsx([self._row(center="CA01", plants="1")]))
        imp.action_validate()
        with self.assertRaises(AccessError):
            imp.line_ids.with_user(self.user_operator).write({"plants": 9.0})

    def test_import_cross_company_center_not_found(self):
        imp = self._make_import(_xlsx([self._row(center="CB01", plants="1")]))
        imp.action_validate()
        self.assertEqual(imp.error_count, 1)
        self.assertIn("no encontrado", (imp.line_ids.error or ""))

    def test_import_cross_company_master_rejected(self):
        unit_b = self.env["step.management.estimation.unit"].create({
            "name": "Bin B", "code": "BINB", "kg_factor": 100.0,
            "company_id": self.company_b.id,
        })
        with self.assertRaises(UserError):
            self._make_import(
                _xlsx([self._row(center="CA01", plants="1")]), unit_id=unit_b.id,
            )

    def test_imported_estimation_flows_into_validation(self):
        imp = self._make_import(
            _xlsx([
                self._row(center="CA01", plants="1000", yield_ue="1"),
                self._row(center="CA02", plants="500", yield_ue="1"),
            ]),
            week_curve_id=self.week_curve.id,
            caliber_curve_id=self.caliber_curve.id,
            class_curve_id=self.class_curve.id,
        )
        imp.action_validate()
        imp.action_import()
        est = imp.estimation_id
        self.assertEqual(est.state, "draft")
        est.action_validate()
        self.assertEqual(est.state, "validated")
        total = est.total_kg
        for axis in ("week", "caliber", "class"):
            axis_lines = est.distribution_ids.filtered(lambda d: d.axis == axis)
            self.assertAlmostEqual(sum(axis_lines.mapped("kg")), total, places=2)
