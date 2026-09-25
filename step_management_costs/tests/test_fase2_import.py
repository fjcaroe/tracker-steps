import base64
import io

import openpyxl

from odoo.exceptions import AccessError, UserError
from odoo.tests import tagged

from .test_management_costs import ManagementCostsCommon, extra_product_vals

HEADERS = [
    "Versión Ppto", "Temporada", "Año", "Mes", "Fundo", "Especie", "Variedad",
    "Tipo CCosto", "Centro de costos", "Origen", "Grupo Presupuesto", "Actividad",
    "Producto-labor", "UdM", "Cantidad", "Valor Ppto$", "TC Ppto", "Valor Ppto US$",
]


def _xlsx(rows, header_at=1, trailing_junk=False):
    wb = openpyxl.Workbook()
    ws = wb.active
    for _ in range(header_at - 1):
        ws.append(["Título del reporte"] + [""] * 17)
    ws.append(HEADERS)
    for r in rows:
        ws.append(r)
    if trailing_junk:
        ws.append([""] * 18)
        ws.append([""] * 18)
        for i in range(20):
            ws.append(["basura residual"] + [""] * 17)
    bio = io.BytesIO()
    wb.save(bio)
    return base64.b64encode(bio.getvalue())


@tagged("post_install", "-at_install")
class TestFase2(ManagementCostsCommon):
    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.group_cost = cls.group_a  # flow_type 'cost' por defecto
        cls.group_income = cls.env["step.management.budget.group"].create({
            "code": "ING", "name": "Ventas", "company_id": cls.company_a.id,
            "flow_type": "income",
        })
        cls.categ = cls.env["product.category"].create({"name": "MC Categ"})
        cls.categ_child = cls.env["product.category"].create({
            "name": "MC Categ hija", "parent_id": cls.categ.id,
        })

    def _row(self, center="CA01", origin="Mano de obra", group="MOA", month=6,
             qty=10, amount=400000, activity="Poda", product="", uom="", year=2026,
             season="2026/2027", version="1"):
        return [version, season, year, month, "", "", "", "Productivo", center,
                origin, group, activity, product, uom, qty, amount, 900, ""]

    # ------------------------------------------------------------------ flow_type
    def test_flow_type_totals_and_margin(self):
        budget = self._new_budget()
        budget.action_generate_lines()
        cost_total = budget.total_amount
        self.assertGreater(cost_total, 0)
        # agrega una línea de ingreso directamente
        self.env["step.management.budget.line"].create({
            "budget_id": budget.id, "center_id": self.center_a.id,
            "group_id": self.group_income.id, "category": "other",
            "indicator": "Venta fruta", "hectares": 10.0,
            "quantity": 1.0, "unit_price": 5000000.0,
            "month_ids": [(0, 0, {"month": "jun", "quantity": 1.0, "unit_price": 5000000.0})],
        })
        budget.invalidate_recordset()
        self.assertAlmostEqual(budget.total_income, 5000000.0)
        self.assertAlmostEqual(budget.total_amount, cost_total)  # ingreso NO suma al costo
        self.assertAlmostEqual(budget.margin, 5000000.0 - cost_total)

    # ------------------------------------------------- resolución de grupo en producto
    def test_product_group_resolution(self):
        Product = self.env["product.product"]
        extra = extra_product_vals(self.env)
        p_own = Product.create(dict({
            "name": "P own", "categ_id": self.categ_child.id,
            "management_budget_group_id": self.group_cost.id,
        }, **extra))
        self.assertEqual(p_own._get_management_budget_group(), self.group_cost)

        self.categ_child.management_budget_group_id = self.group_income.id
        p_cat = Product.create(dict({"name": "P cat", "categ_id": self.categ_child.id}, **extra))
        self.assertEqual(p_cat._get_management_budget_group(), self.group_income)

        self.categ_child.management_budget_group_id = False
        self.categ.management_budget_group_id = self.group_cost.id
        self.assertEqual(p_cat._get_management_budget_group(), self.group_cost)

        self.categ.management_budget_group_id = False
        p_none = Product.create(dict({"name": "P none", "categ_id": self.categ_child.id}, **extra))
        self.assertFalse(p_none._get_management_budget_group())

    # ------------------------------------------------------------------ importación
    def _make_import(self, b64):
        return self.env["step.management.budget.import"].create({
            "company_id": self.company_a.id,
            "file": b64, "filename": "ppto.xlsx",
        })

    def test_import_happy_path(self):
        imp = self._make_import(_xlsx([
            self._row(month=6, qty=10, amount=400000),
            self._row(month=7, qty=5, amount=200000),
            self._row(center="CA02", month=6, qty=3, amount=90000),
        ]))
        imp.action_validate()
        self.assertEqual(imp.error_count, 0, imp.line_ids.mapped("error"))
        self.assertEqual(imp.state, "validated")
        action = imp.action_import()
        self.assertEqual(imp.state, "imported")
        budget = self.env["step.management.operational.budget"].browse(action["res_id"])
        self.assertEqual(budget.origin_type, "import")
        self.assertEqual(budget.import_id, imp)
        # 2 líneas (una por centro), meses cuadran
        self.assertEqual(len(budget.line_ids), 2)
        ca01 = budget.line_ids.filtered(lambda l: l.center_id == self.center_a)
        self.assertAlmostEqual(sum(ca01.month_ids.mapped("quantity")), 15.0)
        self.assertAlmostEqual(budget.total_amount, 690000.0)
        self.assertTrue(all(budget.line_ids.mapped("distribution_complete")))

    def test_import_all_or_nothing(self):
        imp = self._make_import(_xlsx([
            self._row(month=6),
            self._row(center="NO EXISTE", month=7),
        ]))
        imp.action_validate()
        self.assertEqual(imp.error_count, 1)
        budgets_before = self.env["step.management.operational.budget"].search_count([])
        with self.assertRaises(UserError):
            imp.action_import()
        self.assertEqual(
            self.env["step.management.operational.budget"].search_count([]), budgets_before
        )
        # con "sólo válidas" sí importa el resto
        imp.import_valid_only = True
        imp.action_import()
        self.assertEqual(imp.state, "imported")
        self.assertEqual(len(imp.budget_id.line_ids), 1)

    def test_import_idempotent_same_file(self):
        b64 = _xlsx([self._row(month=6)])
        imp1 = self._make_import(b64)
        imp1.action_validate()
        imp1.action_import()
        imp2 = self._make_import(b64)
        imp2.action_validate()
        with self.assertRaises(UserError):
            imp2.action_import()

    def test_import_rejects_formula(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(HEADERS)
        ws.append(self._row(month=6))
        ws["O2"] = "=1+1"  # columna Cantidad (15ª), única fila de datos
        bio = io.BytesIO()
        wb.save(bio)
        imp = self._make_import(base64.b64encode(bio.getvalue()))
        imp.action_validate()
        self.assertEqual(imp.error_count, 1)
        self.assertIn("fórmula", (imp.line_ids.error or "").lower())

    def test_import_stops_at_residual_rows(self):
        imp = self._make_import(_xlsx([self._row(month=6), self._row(month=7)],
                                      trailing_junk=True))
        imp.action_validate()
        self.assertEqual(imp.line_count, 2)

    def test_import_income_group_mismatch(self):
        imp = self._make_import(_xlsx([
            self._row(origin="Ingresos", group="MOA", month=6),  # grupo es de costo
        ]))
        imp.action_validate()
        self.assertEqual(imp.error_count, 1)
        self.assertIn("Ingreso", imp.line_ids.error or "")

    def test_import_header_not_first_row(self):
        imp = self._make_import(_xlsx([self._row(month=6)], header_at=4))
        imp.action_validate()
        self.assertEqual(imp.error_count, 0)
        self.assertEqual(imp.line_count, 1)

    def test_duplicate_rows_in_same_month_are_aggregated(self):
        imp = self._make_import(_xlsx([
            self._row(month=6, qty=2, amount=2000),
            self._row(month=6, qty=3, amount=3000),
        ]))
        imp.action_validate()
        imp.action_import()
        line = imp.budget_id.line_ids
        self.assertEqual(len(line.month_ids), 1)
        self.assertAlmostEqual(line.month_ids.quantity, 5.0)
        self.assertAlmostEqual(line.month_ids.amount, 5000.0)

    def test_operator_cannot_tamper_with_validated_staging(self):
        imp = self._make_import(_xlsx([self._row(month=6)]))
        imp.action_validate()
        with self.assertRaises(AccessError):
            imp.line_ids.with_user(self.user_operator).write({"amount": 1.0})

    def test_xlsm_is_rejected(self):
        imp = self.env["step.management.budget.import"].create({
            "company_id": self.company_a.id,
            "file": _xlsx([self._row(month=6)]),
            "filename": "ppto.xlsm",
        })
        with self.assertRaises(UserError):
            imp.action_validate()
