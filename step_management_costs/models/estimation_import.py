"""Carga normalizada de estimación de cosecha desde Excel (Anexo 1.6.4.1) con
staging.

Reglas (plan D10, iguales que `budget_import.py`):
- límite de filas; no se recorren dimensiones residuales del libro;
- sólo `.xlsx`, sin fórmulas ni contenido activo como fuente de datos;
- los maestros se resuelven por clave única de empresa, nunca por coincidencia
  vaga;
- vista previa y errores por fila antes de confirmar;
- hash de archivo/línea para idempotencia; el mismo archivo no se reimporta;
- todo-o-nada por lote, salvo «importar sólo filas válidas».

La carga **nunca** produce una estimación validada: crea una
`step.management.estimation` en borrador que el operador revisa y el aprobador
valida con las tres curvas (corte 2).
"""

import base64
import hashlib
import io
import re
import unicodedata

from odoo import api, fields, models, _
from odoo.exceptions import UserError, ValidationError

try:
    import openpyxl
except ImportError:  # pragma: no cover - openpyxl es dependencia de Odoo
    openpyxl = None

MAX_DATA_ROWS = 5000
MAX_FILE_BYTES = 10 * 1024 * 1024

# Cabecera esperada, normalizada (minúsculas, sin acentos, sin espacios extra).
EXPECTED_HEADERS = {
    "centro de costo": "center",
    "centro de costos": "center",
    "codigo centro": "center",
    "codigo centro de costo": "center",
    "cuartel": "center",
    "hectareas": "hectares",
    "has": "hectares",
    "ha": "hectares",
    "superficie": "hectares",
    "plantas": "plants",
    "n plantas": "plants",
    "n de plantas": "plants",
    "numero de plantas": "plants",
    "rendimiento ue": "yield_ue",
    "rendimiento": "yield_ue",
    "rendimiento (ue)": "yield_ue",
    "rendimiento por planta": "yield_ue",
    "rendimiento por hectarea": "yield_ue",
    "kilos": "total_kg",
    "total kg": "total_kg",
    "kg": "total_kg",
    "kilos estimados": "total_kg",
    "temporada": "season",
    "fundo": "farm",
    "especie": "species",
    "variedad": "variety",
}
REQUIRED_LOGICAL_COLS = {"center"}


def _norm(value):
    if value is None:
        return ""
    text = str(value).strip().lower()
    text = "".join(
        ch for ch in unicodedata.normalize("NFKD", text) if not unicodedata.combining(ch)
    )
    return re.sub(r"\s+", " ", text)


def _to_float(value):
    if value in (None, ""):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace(" ", "")
    if "," in text and "." in text:
        text = text.replace(".", "").replace(",", ".")
    elif "," in text:
        text = text.replace(",", ".")
    try:
        return float(text)
    except ValueError:
        return None


class StepManagementEstimationImport(models.Model):
    _name = "step.management.estimation.import"
    _description = "Carga de estimación de cosecha desde Excel"
    _inherit = ["mail.thread"]
    _order = "create_date desc, id desc"
    _check_company_auto = True

    name = fields.Char(
        string="Folio", required=True, readonly=True, copy=False,
        default=lambda s: _("Nuevo"),
    )
    company_id = fields.Many2one(
        "res.company", string="Empresa", required=True,
        default=lambda s: s.env.company, index=True,
    )
    state = fields.Selection(
        [("draft", "Borrador"), ("validated", "Validado"),
         ("imported", "Importado"), ("cancelled", "Cancelado")],
        default="draft", required=True, tracking=True, index=True,
    )
    file = fields.Binary(string="Archivo Excel", required=True, attachment=True)
    filename = fields.Char(string="Nombre del archivo")
    file_hash = fields.Char(string="Hash del archivo", size=64, readonly=True, copy=False)
    user_id = fields.Many2one(
        "res.users", string="Cargado por", default=lambda s: s.env.user, readonly=True,
    )

    version_id = fields.Many2one(
        "step.management.estimation.version", string="Versión", required=True,
        check_company=True, domain="[('company_id', '=', company_id)]",
    )
    season = fields.Char(string="Temporada detectada", readonly=True)
    unit_id = fields.Many2one(
        "step.management.estimation.unit", string="Unidad de estimación",
        required=True, check_company=True, domain="[('company_id', '=', company_id)]",
    )
    method = fields.Selection(
        [("plants", "Plantas"), ("hectares", "Hectáreas"), ("kilos", "Kilos")],
        string="Método de cálculo", required=True, default="plants",
    )
    default_yield_ue = fields.Float(
        string="Rendimiento por defecto (UE)", digits=(16, 6),
        help="Se usa cuando una fila no trae rendimiento. Ignorado en el "
             "método «Kilos».",
    )
    species = fields.Char(string="Especie por defecto")
    variety = fields.Char(string="Variedad por defecto")
    week_curve_id = fields.Many2one(
        "step.management.estimation.curve", string="Curva semanal", check_company=True,
        domain="[('company_id', '=', company_id), ('curve_type', '=', 'week'), "
               "('state', '=', 'validated'), ('active', '=', True)]",
    )
    caliber_curve_id = fields.Many2one(
        "step.management.estimation.curve", string="Curva de calibre", check_company=True,
        domain="[('company_id', '=', company_id), ('curve_type', '=', 'caliber'), "
               "('state', '=', 'validated'), ('active', '=', True)]",
    )
    class_curve_id = fields.Many2one(
        "step.management.estimation.curve", string="Curva de clases", check_company=True,
        domain="[('company_id', '=', company_id), ('curve_type', '=', 'class'), "
               "('state', '=', 'validated'), ('active', '=', True)]",
    )

    import_valid_only = fields.Boolean(
        string="Importar sólo filas válidas",
        help="Por omisión la carga es todo-o-nada. Marque esta casilla para "
             "importar únicamente las filas sin error.",
    )
    estimation_id = fields.Many2one(
        "step.management.estimation", string="Estimación creada",
        readonly=True, copy=False,
    )
    line_ids = fields.One2many(
        "step.management.estimation.import.line", "import_id", string="Filas",
    )
    line_count = fields.Integer(compute="_compute_counts", store=True)
    error_count = fields.Integer(compute="_compute_counts", store=True)
    ok_count = fields.Integer(compute="_compute_counts", store=True)
    notes = fields.Text(string="Observaciones")

    @api.depends("line_ids.state")
    def _compute_counts(self):
        for record in self:
            record.line_count = len(record.line_ids)
            record.error_count = len(record.line_ids.filtered(lambda l: l.state == "error"))
            record.ok_count = len(record.line_ids.filtered(lambda l: l.state == "ok"))

    @api.model_create_multi
    def create(self, vals_list):
        for vals in vals_list:
            if vals.get("name", _("Nuevo")) == _("Nuevo"):
                vals["name"] = self.env["ir.sequence"].next_by_code(
                    "step.management.estimation.import"
                ) or _("Nuevo")
        return super().create(vals_list)

    def _reset_to_draft(self):
        self.line_ids.sudo().unlink()
        self.write({"state": "draft", "file_hash": False, "season": False})

    def action_reset(self):
        for record in self:
            if record.state == "imported":
                raise UserError(_("No se puede restablecer una carga ya importada."))
            record._reset_to_draft()

    def action_cancel(self):
        self.filtered(lambda r: r.state != "imported").write({"state": "cancelled"})

    # ------------------------------------------------------------------
    # Validación / parsing
    # ------------------------------------------------------------------
    def _open_workbook(self):
        self.ensure_one()
        if openpyxl is None:
            raise UserError(_("La librería openpyxl no está disponible en el servidor."))
        if not self.file:
            raise UserError(_("Adjunte un archivo Excel."))
        if self.filename and not self.filename.lower().endswith(".xlsx"):
            raise UserError(_("El archivo debe ser .xlsx."))
        try:
            raw = base64.b64decode(self.file, validate=True)
        except (ValueError, TypeError) as exc:
            raise UserError(_("El archivo adjunto no contiene datos Base64 válidos.")) from exc
        if len(raw) > MAX_FILE_BYTES:
            raise UserError(_(
                "El archivo supera el máximo permitido de %(size)s MB."
            ) % {"size": MAX_FILE_BYTES // (1024 * 1024)})
        self.file_hash = hashlib.sha256(raw).hexdigest()
        try:
            return openpyxl.load_workbook(io.BytesIO(raw), data_only=False, read_only=True)
        except Exception as exc:  # noqa: BLE001
            raise UserError(_("No se pudo abrir el Excel: %s") % exc)

    def _locate_header(self, ws):
        for r_idx, row in enumerate(ws.iter_rows(min_row=1, max_row=15, values_only=True), start=1):
            mapping = {}
            for c_idx, value in enumerate(row):
                logical = EXPECTED_HEADERS.get(_norm(value))
                if logical and logical not in mapping:
                    mapping[logical] = c_idx
            if REQUIRED_LOGICAL_COLS.issubset(set(mapping)):
                return r_idx + 1, mapping
        raise UserError(_(
            "No se encontró la fila de cabecera del Anexo 1.6.4.1. Se espera al "
            "menos la columna «Centro de costo / Cuartel»."
        ))

    def action_validate(self):
        self.ensure_one()
        if self.state == "imported":
            raise UserError(_("Esta carga ya fue importada."))
        self.line_ids.sudo().unlink()
        wb = self._open_workbook()
        ws = wb.worksheets[0]
        start_row, cols = self._locate_header(ws)

        def cell(row, key):
            idx = cols.get(key)
            if idx is None or idx >= len(row):
                return None
            return row[idx]

        seasons = set()
        line_vals = []
        empty_streak = 0
        for offset, row in enumerate(
            ws.iter_rows(min_row=start_row, max_row=start_row + MAX_DATA_ROWS - 1, values_only=True)
        ):
            row_number = start_row + offset
            mapped = {key: cell(row, key) for key in cols}
            if all((v is None or str(v).strip() == "") for v in mapped.values()):
                empty_streak += 1
                if empty_streak >= 2:
                    break
                continue
            empty_streak = 0
            vals = self._parse_row(row_number, mapped)
            if vals.get("season"):
                seasons.add(vals["season"].strip())
            line_vals.append(vals)
        wb.close()

        if not line_vals:
            raise UserError(_("El archivo no contiene filas de datos."))

        # centro repetido en el archivo: la 2ª y siguientes van a error
        seen_centers = {}
        for vals in line_vals:
            center_id = vals.get("center_id")
            if not center_id:
                continue
            if center_id in seen_centers:
                extra = _("El centro/cuartel está repetido en el archivo (fila %s).") % seen_centers[center_id]
                vals["error"] = ("%s\n%s" % (vals.get("error") or "", extra)).strip()
                vals["error_field"] = vals["error_field"] or extra[:60]
                vals["state"] = "error"
            else:
                seen_centers[center_id] = vals["row_number"]

        self.write({
            "state": "validated",
            "season": ", ".join(sorted(seasons)) or False,
        })
        self.env["step.management.estimation.import.line"].sudo().create([
            dict(vals, import_id=self.id) for vals in line_vals
        ])
        self.invalidate_recordset(["line_ids", "line_count", "error_count", "ok_count"])
        self.message_post(body=_(
            "Validación: %(total)s fila(s), %(err)s con error."
        ) % {"total": self.line_count, "err": self.error_count})

    def _parse_row(self, row_number, mapped):
        Center = self.env["step.management.cost.center"]
        company = self.company_id

        raw = {
            "raw_center": mapped.get("center"), "raw_hectares": mapped.get("hectares"),
            "raw_plants": mapped.get("plants"), "raw_yield_ue": mapped.get("yield_ue"),
            "raw_total_kg": mapped.get("total_kg"), "raw_season": mapped.get("season"),
            "raw_farm": mapped.get("farm"), "raw_species": mapped.get("species"),
            "raw_variety": mapped.get("variety"),
        }
        vals = {k: ("" if v is None else str(v)) for k, v in raw.items()}
        vals["row_number"] = row_number
        vals["season"] = (str(mapped.get("season")).strip() if mapped.get("season") else "")

        errors = []

        for key, value in mapped.items():
            if isinstance(value, str) and value.strip().startswith("="):
                errors.append(_("La celda «%s» contiene una fórmula.") % key)

        raw_center = (str(mapped.get("center")).strip() if mapped.get("center") else "")
        if not raw_center:
            errors.append(_("Falta el centro de costo / cuartel."))
        else:
            centers = Center.search([
                ("company_id", "=", company.id),
                "|", ("code", "=", raw_center), ("name", "=", raw_center),
            ], limit=2)
            if not centers:
                errors.append(_("Centro/cuartel «%s» no encontrado en la empresa.") % raw_center)
            elif len(centers) > 1:
                errors.append(_("Centro/cuartel «%s» ambiguo.") % raw_center)
            else:
                vals["center_id"] = centers.id

        hectares = _to_float(mapped.get("hectares"))
        if raw["raw_hectares"] not in (None, "") and (hectares is None or hectares < 0):
            errors.append(_("Hectáreas inválidas."))
        elif hectares is not None:
            vals["hectares"] = hectares

        plants = _to_float(mapped.get("plants"))
        if raw["raw_plants"] not in (None, "") and (plants is None or plants < 0):
            errors.append(_("Número de plantas inválido."))
        elif plants is not None:
            vals["plants"] = plants

        yield_ue = _to_float(mapped.get("yield_ue"))
        if raw["raw_yield_ue"] not in (None, "") and (yield_ue is None or yield_ue < 0):
            errors.append(_("Rendimiento inválido."))
        elif yield_ue is not None:
            vals["yield_ue"] = yield_ue

        total_kg = _to_float(mapped.get("total_kg"))
        if raw["raw_total_kg"] not in (None, "") and (total_kg is None or total_kg < 0):
            errors.append(_("Kilos inválidos."))
        elif total_kg is not None:
            vals["total_kg"] = total_kg

        if self.method == "kilos":
            effective_kg = total_kg if total_kg is not None else 0.0
            if effective_kg <= 0:
                errors.append(_("El método «Kilos» exige kilos mayores que cero por fila."))
        else:
            effective_yield = yield_ue if yield_ue is not None else self.default_yield_ue
            if not effective_yield or effective_yield <= 0:
                errors.append(_(
                    "El método «%s» exige un rendimiento mayor que cero "
                    "(en la fila o por defecto)."
                ) % dict(self._fields["method"].selection).get(self.method))

        content = "|".join(_norm(mapped.get(k)) for k in sorted(mapped))
        vals["line_hash"] = hashlib.sha256(
            ("%s|%s|%s" % (self.file_hash or "", row_number, content)).encode("utf-8")
        ).hexdigest()

        vals["state"] = "error" if errors else "ok"
        vals["error"] = "\n".join(errors)
        vals["error_field"] = errors[0][:60] if errors else False
        return vals

    # ------------------------------------------------------------------
    # Importación (todo o nada) -> estimación en BORRADOR
    # ------------------------------------------------------------------
    def action_import(self):
        self.ensure_one()
        if self.state != "validated":
            raise UserError(_("Primero valide el archivo."))
        if self.error_count and not self.import_valid_only:
            raise UserError(_(
                "Hay %s fila(s) con error. Corríjalas y vuelva a validar, o marque "
                "«Importar sólo filas válidas»."
            ) % self.error_count)

        duplicate = self.search([
            ("id", "!=", self.id), ("company_id", "=", self.company_id.id),
            ("file_hash", "=", self.file_hash), ("state", "=", "imported"),
        ], limit=1)
        if duplicate:
            raise UserError(_(
                "Este archivo ya fue importado en %s. No se vuelve a cargar."
            ) % duplicate.name)

        rows = self.line_ids.filtered(lambda l: l.state == "ok")
        if not rows:
            raise UserError(_("No hay filas válidas para importar."))

        seasons = set(rows.mapped("season")) - {False, ""}
        if len(seasons) > 1:
            raise UserError(_(
                "El archivo mezcla temporadas distintas: %s"
            ) % ", ".join(sorted(seasons)))
        season = (list(seasons) or [self.version_id.season])[0]

        estimation = self.env["step.management.estimation"].create({
            "company_id": self.company_id.id,
            "version_id": self.version_id.id,
            "season": season,
            "unit_id": self.unit_id.id,
            "method": self.method,
            "default_yield_ue": self.default_yield_ue,
            "species": self.species or False,
            "variety": self.variety or False,
            "week_curve_id": self.week_curve_id.id or False,
            "caliber_curve_id": self.caliber_curve_id.id or False,
            "class_curve_id": self.class_curve_id.id or False,
            "date": fields.Date.context_today(self),
            "import_id": self.id,
            "center_ids": [(6, 0, rows.center_id.ids)],
        })

        line_cmds = []
        for row in rows:
            center = row.center_id
            has_hectares = bool((row.raw_hectares or "").strip())
            has_plants = bool((row.raw_plants or "").strip())
            has_yield = bool((row.raw_yield_ue or "").strip())
            line_cmds.append((0, 0, {
                "center_id": center.id,
                "farm": (row.raw_farm or "").strip() or center.farm,
                "plot": center.plot,
                "species": (row.raw_species or "").strip() or self.species or center.species,
                "variety": (row.raw_variety or "").strip() or self.variety or center.variety,
                "hectares": row.hectares if has_hectares else center.hectares,
                "plants": row.plants if has_plants else center.plants,
                "yield_ue": row.yield_ue if has_yield else self.default_yield_ue,
                "total_kg_input": row.total_kg if self.method == "kilos" else 0.0,
            }))
        estimation.write({"line_ids": line_cmds})

        self.write({"state": "imported", "estimation_id": estimation.id})
        rows.sudo().write({"imported": True})
        self.message_post(body=_(
            "Importado a %(est)s: %(lines)s línea(s) de %(rows)s fila(s) del Excel. "
            "La estimación queda en borrador para revisión y validación."
        ) % {"est": estimation.name, "lines": len(line_cmds), "rows": len(rows)})
        estimation.message_post(body=_("Creada por importación %s.") % self.name)
        return {
            "type": "ir.actions.act_window",
            "res_model": "step.management.estimation",
            "res_id": estimation.id, "view_mode": "form", "target": "current",
        }


class StepManagementEstimationImportLine(models.Model):
    _name = "step.management.estimation.import.line"
    _description = "Fila de carga de estimación"
    _order = "row_number, id"
    _check_company_auto = True

    import_id = fields.Many2one(
        "step.management.estimation.import", required=True, ondelete="cascade", index=True,
    )
    company_id = fields.Many2one(related="import_id.company_id", store=True, index=True)
    row_number = fields.Integer(string="Fila")
    state = fields.Selection([("ok", "OK"), ("error", "Error")], default="ok", index=True)
    error = fields.Text(string="Detalle del error")
    error_field = fields.Char(string="Error")
    imported = fields.Boolean(string="Importada", default=False)
    line_hash = fields.Char(size=64, index=True)

    raw_center = fields.Char(string="Centro (texto)")
    raw_hectares = fields.Char(string="Hectáreas (texto)")
    raw_plants = fields.Char(string="Plantas (texto)")
    raw_yield_ue = fields.Char(string="Rendimiento (texto)")
    raw_total_kg = fields.Char(string="Kilos (texto)")
    raw_season = fields.Char(string="Temporada (texto)")
    raw_farm = fields.Char(string="Fundo (texto)")
    raw_species = fields.Char(string="Especie (texto)")
    raw_variety = fields.Char(string="Variedad (texto)")

    season = fields.Char()
    center_id = fields.Many2one("step.management.cost.center", check_company=True)
    hectares = fields.Float(digits=(16, 4))
    plants = fields.Float(digits=(16, 2))
    yield_ue = fields.Float(digits=(16, 6))
    total_kg = fields.Float(digits=(16, 2))

    @api.constrains("center_id", "company_id")
    def _check_company_consistency(self):
        for record in self:
            center = record.center_id
            if center and center.company_id and center.company_id != record.company_id:
                raise ValidationError(_("Fila %s: centro de otra empresa.") % record.row_number)
