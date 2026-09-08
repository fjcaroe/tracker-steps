"""Carga normalizada de presupuesto desde Excel (Anexo 1.6.2.1) con staging.

Reglas (plan D10):
- límite de filas/celdas: no se recorren dimensiones residuales del libro;
- se verifica extensión, cabeceras y no se aceptan fórmulas;
- los maestros se resuelven por claves únicas de empresa, nunca por
  coincidencia vaga;
- vista previa y errores por fila antes de confirmar;
- clave natural + hash del archivo/línea para ser idempotente;
- todo o nada por lote, salvo que el usuario marque explícitamente
  «importar sólo filas válidas».
"""

import base64
import hashlib
import io
import re
import unicodedata

from odoo import api, fields, models, _
from odoo.exceptions import UserError, ValidationError

try:
    import openpyxl
except ImportError:  # pragma: no cover - openpyxl es dependencia de Odoo
    openpyxl = None

MAX_DATA_ROWS = 5000
MAX_FILE_BYTES = 10 * 1024 * 1024

MONTH_KEY_BY_NUMBER = {
    1: "jan", 2: "feb", 3: "mar", 4: "apr", 5: "may", 6: "jun",
    7: "jul", 8: "aug", 9: "sep", 10: "oct", 11: "nov", 12: "dec",
}
MONTH_SELECTION = [
    ("may", "Mayo"), ("jun", "Junio"), ("jul", "Julio"), ("aug", "Agosto"),
    ("sep", "Septiembre"), ("oct", "Octubre"), ("nov", "Noviembre"),
    ("dec", "Diciembre"), ("jan", "Enero"), ("feb", "Febrero"),
    ("mar", "Marzo"), ("apr", "Abril"),
]

ORIGIN_TO_CATEGORY = {
    "mano de obra": "labor",
    "insumos": "input",
    "insumo agricola": "input",
    "insumos agricolas": "input",
    "maquinaria": "machinery",
    "maquinarias": "machinery",
    "gastos y servicios": "service",
    "servicio": "service",
    "servicios": "service",
    "ingresos": "other",
    "ingreso": "other",
    "otro": "other",
}
INCOME_ORIGINS = {"ingresos", "ingreso"}

# Cabecera esperada, normalizada (minúsculas, sin acentos, sin espacios extra).
EXPECTED_HEADERS = {
    "version ppto": "version",
    "temporada": "season",
    "ano": "year",
    "mes": "month",
    "fundo": "farm",
    "especie": "species",
    "variedad": "variety",
    "tipo ccosto": "cost_type",
    "centro de costos": "center",
    "origen": "origin",
    "grupo presupuesto": "group",
    "actividad": "activity",
    "producto-labor": "product",
    "producto labor": "product",
    "udm": "uom",
    "cantidad": "quantity",
    "valor ppto$": "amount",
    "valor ppto $": "amount",
    "tc ppto": "rate",
    "valor ppto us$": "amount_usd",
    "valor ppto us $": "amount_usd",
}
REQUIRED_LOGICAL_COLS = {
    "season", "year", "month", "center", "origin", "group", "quantity", "amount",
}


def _norm(value):
    if value is None:
        return ""
    text = str(value).strip().lower()
    text = "".join(
        ch for ch in unicodedata.normalize("NFKD", text) if not unicodedata.combining(ch)
    )
    return re.sub(r"\s+", " ", text)


def _to_float(value):
    if value in (None, ""):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace(" ", "")
    # admite 1.234.567,89 y 1234567.89
    if "," in text and "." in text:
        text = text.replace(".", "").replace(",", ".")
    elif "," in text:
        text = text.replace(",", ".")
    try:
        return float(text)
    except ValueError:
        return None


class StepManagementBudgetImport(models.Model):
    _name = "step.management.budget.import"
    _description = "Carga de presupuesto desde Excel"
    _inherit = ["mail.thread"]
    _order = "create_date desc, id desc"
    _check_company_auto = True

    name = fields.Char(string="Folio", required=True, readonly=True, copy=False, default=lambda s: _("Nuevo"))
    company_id = fields.Many2one(
        "res.company", string="Empresa", required=True, default=lambda s: s.env.company, index=True,
    )
    currency_id = fields.Many2one(related="company_id.currency_id")
    state = fields.Selection(
        [("draft", "Borrador"), ("validated", "Validado"),
         ("imported", "Importado"), ("cancelled", "Cancelado")],
        default="draft", required=True, tracking=True, index=True,
    )
    file = fields.Binary(string="Archivo Excel", required=True, attachment=True)
    filename = fields.Char(string="Nombre del archivo")
    file_hash = fields.Char(string="Hash del archivo", size=64, readonly=True, copy=False)
    user_id = fields.Many2one("res.users", string="Cargado por", default=lambda s: s.env.user, readonly=True)
    season = fields.Char(string="Temporada detectada", readonly=True)
    budget_version = fields.Char(string="Versión detectada", readonly=True)
    import_valid_only = fields.Boolean(
        string="Importar sólo filas válidas",
        help="Por omisión la carga es todo-o-nada. Marque esta casilla para "
             "importar únicamente las filas sin error.",
    )
    budget_id = fields.Many2one(
        "step.management.operational.budget", string="Presupuesto creado", readonly=True, copy=False,
    )
    line_ids = fields.One2many("step.management.budget.import.line", "import_id", string="Filas")
    line_count = fields.Integer(compute="_compute_counts", store=True)
    error_count = fields.Integer(compute="_compute_counts", store=True)
    ok_count = fields.Integer(compute="_compute_counts", store=True)
    notes = fields.Text(string="Observaciones")

    @api.depends("line_ids.state")
    def _compute_counts(self):
        for record in self:
            record.line_count = len(record.line_ids)
            record.error_count = len(record.line_ids.filtered(lambda l: l.state == "error"))
            record.ok_count = len(record.line_ids.filtered(lambda l: l.state == "ok"))

    @api.model_create_multi
    def create(self, vals_list):
        for vals in vals_list:
            if vals.get("name", _("Nuevo")) == _("Nuevo"):
                vals["name"] = self.env["ir.sequence"].next_by_code(
                    "step.management.budget.import"
                ) or _("Nuevo")
        return super().create(vals_list)

    def _reset_to_draft(self):
        self.line_ids.sudo().unlink()
        self.write({
            "state": "draft", "file_hash": False, "season": False, "budget_version": False,
        })

    def action_reset(self):
        for record in self:
            if record.state == "imported":
                raise UserError(_("No se puede restablecer una carga ya importada."))
            record._reset_to_draft()

    def action_cancel(self):
        self.filtered(lambda r: r.state != "imported").write({"state": "cancelled"})

    # ------------------------------------------------------------------
    # Validación / parsing
    # ------------------------------------------------------------------
    def _open_workbook(self):
        self.ensure_one()
        if openpyxl is None:
            raise UserError(_("La librería openpyxl no está disponible en el servidor."))
        if not self.file:
            raise UserError(_("Adjunte un archivo Excel."))
        if self.filename and not self.filename.lower().endswith(".xlsx"):
            raise UserError(_("El archivo debe ser .xlsx."))
        try:
            raw = base64.b64decode(self.file, validate=True)
        except (ValueError, TypeError) as exc:
            raise UserError(_("El archivo adjunto no contiene datos Base64 válidos.")) from exc
        if len(raw) > MAX_FILE_BYTES:
            raise UserError(_(
                "El archivo supera el máximo permitido de %(size)s MB."
            ) % {"size": MAX_FILE_BYTES // (1024 * 1024)})
        self.file_hash = hashlib.sha256(raw).hexdigest()
        try:
            # data_only=False: se leen los valores literales de las celdas. Una
            # celda con fórmula devuelve el texto "=..." y se rechaza (D10: no
            # se aceptan fórmulas ni contenido activo como fuente de datos).
            return openpyxl.load_workbook(io.BytesIO(raw), data_only=False, read_only=True)
        except Exception as exc:  # noqa: BLE001
            raise UserError(_("No se pudo abrir el Excel: %s") % exc)

    def _locate_header(self, ws):
        """Devuelve (fila_datos_inicio, {logical_col: idx}). Busca la cabecera en
        las primeras 15 filas por coincidencia de nombres normalizados."""
        for r_idx, row in enumerate(ws.iter_rows(min_row=1, max_row=15, values_only=True), start=1):
            mapping = {}
            for c_idx, value in enumerate(row):
                logical = EXPECTED_HEADERS.get(_norm(value))
                if logical and logical not in mapping:
                    mapping[logical] = c_idx
            if REQUIRED_LOGICAL_COLS.issubset(set(mapping)):
                return r_idx + 1, mapping
        raise UserError(_(
            "No se encontró la fila de cabecera del Anexo 1.6.2.1. Se esperan al "
            "menos las columnas: Temporada, Año, Mes, Centro de costos, Origen, "
            "Grupo Presupuesto, Cantidad, Valor Ppto$."
        ))

    def action_validate(self):
        self.ensure_one()
        if self.state == "imported":
            raise UserError(_("Esta carga ya fue importada."))
        self.line_ids.sudo().unlink()
        wb = self._open_workbook()
        ws = wb.worksheets[0]
        start_row, cols = self._locate_header(ws)

        def cell(row, key):
            idx = cols.get(key)
            if idx is None or idx >= len(row):
                return None
            return row[idx]

        seasons, versions = set(), set()
        line_vals = []
        empty_streak = 0
        for offset, row in enumerate(
            ws.iter_rows(min_row=start_row, max_row=start_row + MAX_DATA_ROWS - 1, values_only=True)
        ):
            row_number = start_row + offset
            mapped = {key: cell(row, key) for key in cols}
            if all((v is None or str(v).strip() == "") for v in mapped.values()):
                empty_streak += 1
                if empty_streak >= 2:
                    break
                continue
            empty_streak = 0
            vals = self._parse_row(row_number, mapped)
            if vals.get("season"):
                seasons.add(vals["season"].strip())
            if vals.get("raw_version"):
                versions.add(str(vals["raw_version"]).strip())
            line_vals.append(vals)
        wb.close()

        if not line_vals:
            raise UserError(_("El archivo no contiene filas de datos."))
        self.write({
            "state": "validated",
            "season": ", ".join(sorted(seasons)) or False,
            "budget_version": ", ".join(sorted(versions)) or False,
        })
        self.env["step.management.budget.import.line"].sudo().create([
            dict(vals, import_id=self.id) for vals in line_vals
        ])
        self.invalidate_recordset(["line_ids", "line_count", "error_count", "ok_count"])
        self.message_post(body=_(
            "Validación: %(total)s fila(s), %(err)s con error."
        ) % {"total": self.line_count, "err": self.error_count})

    def _parse_row(self, row_number, mapped):
        Center = self.env["step.management.cost.center"]
        Group = self.env["step.management.budget.group"]
        Product = self.env["product.product"]
        Uom = self.env["uom.uom"]
        company = self.company_id

        raw = {
            "raw_version": mapped.get("version"), "raw_season": mapped.get("season"),
            "raw_year": mapped.get("year"), "raw_month": mapped.get("month"),
            "raw_farm": mapped.get("farm"), "raw_species": mapped.get("species"),
            "raw_variety": mapped.get("variety"), "raw_cost_type": mapped.get("cost_type"),
            "raw_center": mapped.get("center"), "raw_origin": mapped.get("origin"),
            "raw_group": mapped.get("group"), "raw_activity": mapped.get("activity"),
            "raw_product": mapped.get("product"), "raw_uom": mapped.get("uom"),
            "raw_quantity": mapped.get("quantity"), "raw_amount": mapped.get("amount"),
            "raw_rate": mapped.get("rate"), "raw_amount_usd": mapped.get("amount_usd"),
        }
        vals = {k: ("" if v is None else str(v)) for k, v in raw.items()}
        vals["row_number"] = row_number
        vals["season"] = (str(mapped.get("season")).strip() if mapped.get("season") else "")

        errors = []

        # fórmulas / contenido activo
        for k, v in mapped.items():
            if isinstance(v, str) and v.strip().startswith("="):
                errors.append(_("La celda «%s» contiene una fórmula.") % k)

        year = _to_float(mapped.get("year"))
        if not year or not (1900 <= int(year) <= 2100):
            errors.append(_("Año inválido."))
        else:
            vals["year"] = int(year)

        month_num = _to_float(mapped.get("month"))
        if not month_num or int(month_num) not in MONTH_KEY_BY_NUMBER:
            errors.append(_("Mes inválido (se espera 1-12)."))
        else:
            vals["month"] = MONTH_KEY_BY_NUMBER[int(month_num)]

        raw_center = (str(mapped.get("center")).strip() if mapped.get("center") else "")
        if not raw_center:
            errors.append(_("Falta el centro de costos."))
        else:
            centers = Center.search([
                ("company_id", "=", company.id),
                "|", ("code", "=", raw_center), ("name", "=", raw_center),
            ], limit=2)
            if not centers:
                errors.append(_("Centro de costos «%s» no encontrado en la empresa.") % raw_center)
            elif len(centers) > 1:
                errors.append(_("Centro de costos «%s» ambiguo.") % raw_center)
            else:
                vals["center_id"] = centers.id

        origin_norm = _norm(mapped.get("origin"))
        category = ORIGIN_TO_CATEGORY.get(origin_norm)
        if not category:
            errors.append(_("Origen «%s» no reconocido.") % (mapped.get("origin") or ""))
        else:
            vals["category"] = category
        wants_income = origin_norm in INCOME_ORIGINS

        raw_group = (str(mapped.get("group")).strip() if mapped.get("group") else "")
        if not raw_group:
            errors.append(_("Falta el grupo presupuesto."))
        else:
            groups = Group.search([
                ("company_id", "=", company.id),
                "|", ("code", "=", raw_group), ("name", "=", raw_group),
            ], limit=2)
            if not groups:
                errors.append(_("Grupo presupuesto «%s» no encontrado.") % raw_group)
            elif len(groups) > 1:
                errors.append(_("Grupo presupuesto «%s» ambiguo.") % raw_group)
            else:
                vals["group_id"] = groups.id
                if wants_income and groups.flow_type != "income":
                    errors.append(_("El grupo «%s» no es de tipo Ingreso.") % raw_group)
                if not wants_income and groups.flow_type == "income":
                    errors.append(_("El grupo «%s» es de tipo Ingreso pero el origen no.") % raw_group)

        raw_product = (str(mapped.get("product")).strip() if mapped.get("product") else "")
        if raw_product:
            products = Product.search([
                ("company_id", "in", [False, company.id]),
                "|", ("default_code", "=", raw_product), ("name", "=", raw_product),
            ], limit=2)
            if not products:
                errors.append(_("Producto «%s» no encontrado.") % raw_product)
            elif len(products) > 1:
                errors.append(_("Producto «%s» ambiguo.") % raw_product)
            else:
                vals["product_id"] = products.id

        raw_uom = (str(mapped.get("uom")).strip() if mapped.get("uom") else "")
        if raw_uom:
            uoms = Uom.search([("name", "=", raw_uom)], limit=2)
            if not uoms:
                errors.append(_("Unidad de medida «%s» no encontrada.") % raw_uom)
            elif len(uoms) > 1:
                errors.append(_("Unidad de medida «%s» ambigua.") % raw_uom)
            else:
                vals["uom_id"] = uoms.id

        quantity = _to_float(mapped.get("quantity"))
        if quantity is None or quantity < 0:
            errors.append(_("Cantidad inválida."))
        else:
            vals["quantity"] = quantity

        amount = _to_float(mapped.get("amount"))
        if amount is None:
            errors.append(_("Valor Ppto$ inválido."))
        else:
            vals["amount"] = amount
            if quantity == 0 and amount:
                errors.append(_(
                    "Una fila agrícola con cantidad cero no puede tener monto. "
                    "Use un presupuesto general en modo «Monto directo»."
                ))

        rate = _to_float(mapped.get("rate"))
        vals["rate"] = rate or 0.0

        # hash de línea: archivo + fila + contenido normalizado de las columnas mapeadas
        content = "|".join(_norm(mapped.get(k)) for k in sorted(mapped))
        vals["line_hash"] = hashlib.sha256(
            ("%s|%s|%s" % (self.file_hash or "", row_number, content)).encode("utf-8")
        ).hexdigest()

        vals["state"] = "error" if errors else "ok"
        vals["error"] = "\n".join(errors)
        vals["error_field"] = errors[0][:60] if errors else False
        return vals

    # ------------------------------------------------------------------
    # Importación (todo o nada)
    # ------------------------------------------------------------------
    def action_import(self):
        self.ensure_one()
        if self.state != "validated":
            raise UserError(_("Primero valide el archivo."))
        if self.error_count and not self.import_valid_only:
            raise UserError(_(
                "Hay %s fila(s) con error. Corríjalas y vuelva a validar, o marque "
                "«Importar sólo filas válidas»."
            ) % self.error_count)

        duplicate = self.search([
            ("id", "!=", self.id), ("company_id", "=", self.company_id.id),
            ("file_hash", "=", self.file_hash), ("state", "=", "imported"),
        ], limit=1)
        if duplicate:
            raise UserError(_(
                "Este archivo ya fue importado en %s. No se vuelve a cargar."
            ) % duplicate.name)

        rows = self.line_ids.filtered(lambda l: l.state == "ok")
        if not rows:
            raise UserError(_("No hay filas válidas para importar."))

        seasons = set(rows.mapped("season")) - {False, ""}
        if len(seasons) > 1:
            raise UserError(_("El archivo mezcla temporadas distintas: %s") % ", ".join(sorted(seasons)))
        versions = set(r.raw_version.strip() for r in rows if r.raw_version) or {""}
        if len(versions) > 1:
            raise UserError(_("El archivo mezcla versiones distintas: %s") % ", ".join(sorted(versions)))

        budget = self.env["step.management.operational.budget"].create({
            "description": _("Importación %s") % self.name,
            "company_id": self.company_id.id,
            "season": (list(seasons) or [_("Sin temporada")])[0],
            "budget_version": (list(versions) or [""])[0] or False,
            "date": fields.Date.context_today(self),
            "import_id": self.id,
        })

        # agrupar filas por línea de presupuesto
        buckets = {}
        for row in rows:
            key = (
                row.center_id.id, row.group_id.id, row.category or "other",
                row.product_id.id, (row.raw_activity or "").strip(),
                (row.raw_product or row.raw_activity or _("Importado")).strip(),
                row.uom_id.id,
            )
            buckets.setdefault(key, []).append(row)

        allocation_seen = set()
        alloc_cmds, line_cmds = [], []
        for key, bucket_rows in buckets.items():
            center_id, group_id, category, product_id, activity, indicator, uom_id = key
            center = self.env["step.management.cost.center"].browse(center_id)
            if center_id not in allocation_seen:
                allocation_seen.add(center_id)
                alloc_cmds.append((0, 0, {
                    "center_id": center_id,
                    "hectares": center.hectares if center.hectares > 0 else 1.0,
                }))
            total_qty = sum(r.quantity for r in bucket_rows)
            total_amt = sum(r.amount for r in bucket_rows)
            unit_price = (total_amt / total_qty) if total_qty else 0.0
            hectares = center.hectares if center.hectares > 0 else 1.0
            month_buckets = {}
            for r in bucket_rows:
                month_bucket = month_buckets.setdefault(
                    r.month, {"quantity": 0.0, "amount": 0.0}
                )
                month_bucket["quantity"] += r.quantity
                month_bucket["amount"] += r.amount
            month_cmds = []
            for month, values in month_buckets.items():
                m_price = (
                    values["amount"] / values["quantity"]
                    if values["quantity"] else 0.0
                )
                month_cmds.append((0, 0, {
                    "month": month, "quantity": values["quantity"],
                    "unit_price": m_price,
                }))
            line_cmds.append((0, 0, {
                "center_id": center_id, "group_id": group_id, "category": category,
                "indicator": indicator or _("Importado"), "activity": activity or False,
                "product_id": product_id or False, "uom_id": uom_id or False,
                "hectares": hectares, "quantity_per_ha": (total_qty / hectares) if hectares else 0.0,
                "quantity": total_qty, "unit_price": unit_price,
                "calculation_mode": "quantity",
                "month_ids": month_cmds,
            }))

        budget.write({
            "allocation_ids": alloc_cmds, "line_ids": line_cmds, "state": "calculated",
        })
        self.write({"state": "imported", "budget_id": budget.id})
        rows.sudo().write({"imported": True})
        self.message_post(body=_(
            "Importado a %(budget)s: %(lines)s línea(s), %(rows)s fila(s) del Excel."
        ) % {"budget": budget.name, "lines": len(line_cmds), "rows": len(rows)})
        budget.message_post(body=_("Creado por importación %s.") % self.name)
        return {
            "type": "ir.actions.act_window",
            "res_model": "step.management.operational.budget",
            "res_id": budget.id, "view_mode": "form", "target": "current",
        }


class StepManagementBudgetImportLine(models.Model):
    _name = "step.management.budget.import.line"
    _description = "Fila de carga de presupuesto"
    _order = "row_number, id"
    _check_company_auto = True

    import_id = fields.Many2one("step.management.budget.import", required=True, ondelete="cascade", index=True)
    company_id = fields.Many2one(related="import_id.company_id", store=True, index=True)
    currency_id = fields.Many2one(related="import_id.company_id.currency_id")
    row_number = fields.Integer(string="Fila")
    state = fields.Selection([("ok", "OK"), ("error", "Error")], default="ok", index=True)
    error = fields.Text(string="Detalle del error")
    error_field = fields.Char(string="Error")
    imported = fields.Boolean(string="Importada", default=False)
    line_hash = fields.Char(size=64, index=True)

    raw_version = fields.Char()
    raw_season = fields.Char()
    raw_year = fields.Char()
    raw_month = fields.Char()
    raw_farm = fields.Char()
    raw_species = fields.Char()
    raw_variety = fields.Char()
    raw_cost_type = fields.Char()
    raw_center = fields.Char(string="Centro (texto)")
    raw_origin = fields.Char(string="Origen (texto)")
    raw_group = fields.Char(string="Grupo (texto)")
    raw_activity = fields.Char(string="Actividad (texto)")
    raw_product = fields.Char(string="Producto (texto)")
    raw_uom = fields.Char(string="UdM (texto)")
    raw_quantity = fields.Char(string="Cantidad (texto)")
    raw_amount = fields.Char(string="Valor$ (texto)")
    raw_rate = fields.Char(string="TC (texto)")
    raw_amount_usd = fields.Char()

    season = fields.Char()
    year = fields.Integer()
    month = fields.Selection(MONTH_SELECTION)
    center_id = fields.Many2one("step.management.cost.center", check_company=True)
    group_id = fields.Many2one("step.management.budget.group", check_company=True)
    category = fields.Selection(
        [("labor", "Mano de obra"), ("input", "Insumo agrícola"),
         ("machinery", "Maquinaria"), ("service", "Servicio"), ("other", "Otro")],
    )
    product_id = fields.Many2one("product.product")
    uom_id = fields.Many2one("uom.uom")
    quantity = fields.Float(digits=(16, 4))
    amount = fields.Monetary(currency_field="currency_id")
    rate = fields.Float(digits=(16, 6))

    @api.constrains("center_id", "group_id", "company_id")
    def _check_company_consistency(self):
        for record in self:
            for rel in (record.center_id, record.group_id):
                if rel and rel.company_id and rel.company_id != record.company_id:
                    raise ValidationError(_("Fila %s: maestro de otra empresa.") % record.row_number)
