"""Corte 1 post Fase 6 — carga de la receta de un programa fito/ferti desde
Excel, con el mismo patrón de staging seguro que `budget_import.py` (plan
D10): límite de filas/bytes, sin fórmulas, cabeceras normalizadas, vista
previa con errores por fila, hash de archivo para idempotencia, todo-o-nada
salvo «importar sólo filas válidas».

El encabezado del programa (tipo, temporada, especie/variedad, política de
precio, moneda, centros de costo) se define en el propio registro de carga
ANTES de validar — el Excel sólo aporta las filas de receta (producto,
dosis/ha, UdM, objetivo, semana, carencia, reingreso). Importar crea el
`step.management.crop.program` en **borrador**: nunca lo aprueba
automáticamente («aprobación segura», respuesta del cliente).
"""

import base64
import hashlib
import io
import re
import unicodedata

from psycopg2 import IntegrityError

from odoo import api, fields, models, _
from odoo.exceptions import UserError, ValidationError

try:
    import openpyxl
except ImportError:  # pragma: no cover - openpyxl es dependencia de Odoo
    openpyxl = None

from .crop_program import (
    PROGRAM_TYPES, PRICE_POLICIES, check_center_variety_coherence,
)

MAX_DATA_ROWS = 5000
MAX_FILE_BYTES = 10 * 1024 * 1024

EXPECTED_HEADERS = {
    "producto": "product",
    "dosis ha": "dose_per_ha",
    "dosis/ha": "dose_per_ha",
    "dosis por hectarea": "dose_per_ha",
    "udm": "uom",
    "objetivo": "target",
    "semana": "week_number",
    "carencia": "phi_days",
    "carencia dias": "phi_days",
    "reingreso": "rei_hours",
    "reingreso horas": "rei_hours",
}
REQUIRED_LOGICAL_COLS = {"product", "dose_per_ha"}


def _norm(value):
    if value is None:
        return ""
    text = str(value).strip().lower()
    text = "".join(
        ch for ch in unicodedata.normalize("NFKD", text) if not unicodedata.combining(ch)
    )
    return re.sub(r"\s+", " ", text)


def _to_int_strict(value):
    """(entero, error) — exige que no haya parte decimal (R5: no truncar
    silenciosamente `20.5`, `7.5` o `12.5`). `(None, None)` si `value` está
    vacío (campo opcional); `(None, "decimal")` si tiene parte fraccionaria."""
    num = _to_float(value)
    if num is None:
        return None, None
    if not float(num).is_integer():
        return None, "decimal"
    return int(num), None


def _to_float(value):
    if value in (None, ""):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace(" ", "")
    if "," in text and "." in text:
        text = text.replace(".", "").replace(",", ".")
    elif "," in text:
        text = text.replace(",", ".")
    try:
        return float(text)
    except ValueError:
        return None


class StepManagementCropProgramImport(models.Model):
    _name = "step.management.crop.program.import"
    _description = "Carga de receta de programa fito/ferti desde Excel"
    _inherit = ["mail.thread"]
    _order = "create_date desc, id desc"
    _check_company_auto = True

    name = fields.Char(
        string="Folio", required=True, readonly=True, copy=False,
        default=lambda s: _("Nuevo"),
    )
    company_id = fields.Many2one(
        "res.company", string="Empresa", required=True,
        default=lambda s: s.env.company, index=True,
    )
    currency_id = fields.Many2one(
        "res.currency", string="Moneda", required=True,
        default=lambda s: s.env.company.currency_id,
    )
    state = fields.Selection(
        [("draft", "Borrador"), ("validated", "Validado"),
         ("imported", "Importado"), ("cancelled", "Cancelado")],
        default="draft", required=True, tracking=True, index=True,
    )
    program_type = fields.Selection(
        PROGRAM_TYPES, string="Tipo de programa", required=True, default="phyto",
    )
    season = fields.Char(string="Temporada", required=True, help="Ej.: 2026/2027")
    species = fields.Char(string="Especie")
    variety = fields.Char(string="Variedad")
    date = fields.Date(string="Fecha", default=fields.Date.context_today)
    price_policy = fields.Selection(
        PRICE_POLICIES, string="Política de precio", required=True, default="standard",
    )
    center_ids = fields.Many2many(
        "step.management.cost.center", "step_management_crop_program_import_center_rel",
        "import_id", "center_id", string="Centros de costo",
        domain="[('company_id', '=', company_id)]",
        help="Centros de costo de la temporada; todos deben compartir la "
             "misma variedad, igual que en el programa.",
    )
    file = fields.Binary(string="Archivo Excel", required=True, attachment=True)
    filename = fields.Char(string="Nombre del archivo")
    file_hash = fields.Char(string="Hash del archivo", size=64, readonly=True, copy=False)
    import_key_hash = fields.Char(
        string="Huella de la intención", size=64, readonly=True, copy=False,
        help="R5: hash de empresa + temporada + tipo + centros (normalizados) "
             "+ hash del archivo. La idempotencia se protege sobre esta clave "
             "completa (no sólo el archivo), para permitir reutilizar una "
             "receta legítima en otra temporada o conjunto de centros; un "
             "índice único parcial de base de datos la protege frente a "
             "concurrencia (ver `init()`).",
    )
    user_id = fields.Many2one("res.users", string="Cargado por", default=lambda s: s.env.user, readonly=True)
    import_valid_only = fields.Boolean(
        string="Importar sólo filas válidas",
        help="Por omisión la carga es todo-o-nada. Marque esta casilla para "
             "importar únicamente las filas sin error.",
    )
    program_id = fields.Many2one(
        "step.management.crop.program", string="Programa creado", readonly=True, copy=False,
    )
    line_ids = fields.One2many(
        "step.management.crop.program.import.line", "import_id", string="Filas",
    )
    line_count = fields.Integer(compute="_compute_counts", store=True)
    error_count = fields.Integer(compute="_compute_counts", store=True)
    ok_count = fields.Integer(compute="_compute_counts", store=True)
    notes = fields.Text(string="Observaciones")

    @api.depends("line_ids.state")
    def _compute_counts(self):
        for record in self:
            record.line_count = len(record.line_ids)
            record.error_count = len(record.line_ids.filtered(lambda l: l.state == "error"))
            record.ok_count = len(record.line_ids.filtered(lambda l: l.state == "ok"))

    def init(self):
        # R5: protege la idempotencia frente a concurrencia con un índice
        # único parcial de Postgres (mismo patrón que usa el propio Odoo
        # core, p. ej. `stock.quant`) — no basta el `search()` previo en
        # `action_import()`, que es vulnerable a una carrera entre dos
        # importaciones simultáneas con la misma intención.
        self.env.cr.execute("""
            CREATE UNIQUE INDEX IF NOT EXISTS
                step_management_crop_program_import_key_uniq
            ON step_management_crop_program_import (company_id, import_key_hash)
            WHERE state = 'imported' AND import_key_hash IS NOT NULL
        """)

    @api.model_create_multi
    def create(self, vals_list):
        for vals in vals_list:
            if vals.get("name", _("Nuevo")) == _("Nuevo"):
                vals["name"] = self.env["ir.sequence"].next_by_code(
                    "step.management.crop.program.import"
                ) or _("Nuevo")
        return super().create(vals_list)

    @api.constrains("center_ids", "company_id")
    def _check_center_company(self):
        for record in self:
            wrong = record.center_ids.filtered(
                lambda center: center.company_id != record.company_id
            )
            if wrong:
                raise ValidationError(_(
                    "Los centros de costo deben pertenecer a la empresa "
                    "%(company)s: %(centers)s"
                ) % {
                    "company": record.company_id.display_name,
                    "centers": ", ".join(wrong.mapped("display_name")),
                })

    @api.constrains("center_ids", "variety")
    def _check_center_variety(self):
        # R4: regla compartida con `crop_program.py` (mismo criterio que el
        # programa: mezcla vacío/informado y desacuerdo de encabezado son
        # error, no sólo variedades múltiples).
        for record in self:
            check_center_variety_coherence(record.center_ids, record.variety)

    def _reset_to_draft(self):
        self.line_ids.sudo().unlink()
        self.write({"state": "draft", "file_hash": False})

    def action_reset(self):
        for record in self:
            if record.state == "imported":
                raise UserError(_("No se puede restablecer una carga ya importada."))
            record._reset_to_draft()

    def action_cancel(self):
        self.filtered(lambda r: r.state != "imported").write({"state": "cancelled"})

    # ------------------------------------------------------------------
    # Validación / parsing
    # ------------------------------------------------------------------
    def _open_workbook(self):
        self.ensure_one()
        if openpyxl is None:
            raise UserError(_("La librería openpyxl no está disponible en el servidor."))
        if not self.file:
            raise UserError(_("Adjunte un archivo Excel."))
        if self.filename and not self.filename.lower().endswith(".xlsx"):
            raise UserError(_("El archivo debe ser .xlsx."))
        try:
            raw = base64.b64decode(self.file, validate=True)
        except (ValueError, TypeError) as exc:
            raise UserError(_("El archivo adjunto no contiene datos Base64 válidos.")) from exc
        if len(raw) > MAX_FILE_BYTES:
            raise UserError(_(
                "El archivo supera el máximo permitido de %(size)s MB."
            ) % {"size": MAX_FILE_BYTES // (1024 * 1024)})
        self.file_hash = hashlib.sha256(raw).hexdigest()
        try:
            return openpyxl.load_workbook(io.BytesIO(raw), data_only=False, read_only=True)
        except Exception as exc:  # noqa: BLE001
            raise UserError(_("No se pudo abrir el Excel: %s") % exc)

    def _locate_header(self, ws):
        for r_idx, row in enumerate(ws.iter_rows(min_row=1, max_row=15, values_only=True), start=1):
            mapping = {}
            for c_idx, value in enumerate(row):
                logical = EXPECTED_HEADERS.get(_norm(value))
                if logical and logical not in mapping:
                    mapping[logical] = c_idx
            if REQUIRED_LOGICAL_COLS.issubset(set(mapping)):
                return r_idx + 1, mapping
        raise UserError(_(
            "No se encontró la fila de cabecera de la receta. Se esperan al "
            "menos las columnas: Producto, Dosis/Ha."
        ))

    def action_validate(self):
        self.ensure_one()
        if self.state == "imported":
            raise UserError(_("Esta carga ya fue importada."))
        if not self.center_ids:
            raise UserError(_(
                "Seleccione al menos un centro de costo antes de validar."
            ))
        self.line_ids.sudo().unlink()
        wb = self._open_workbook()
        # R5: la clave de idempotencia representa la intención completa
        # (empresa, temporada, tipo, centros normalizados y archivo), no
        # sólo empresa+archivo — así una receta legítima se puede reutilizar
        # en otra temporada o con otro conjunto de centros.
        self.import_key_hash = hashlib.sha256("|".join([
            str(self.company_id.id), self.season or "", self.program_type,
            ",".join(sorted(self.center_ids.mapped("code"))), self.file_hash or "",
        ]).encode("utf-8")).hexdigest()
        ws = wb.worksheets[0]
        start_row, cols = self._locate_header(ws)

        def cell(row, key):
            idx = cols.get(key)
            if idx is None or idx >= len(row):
                return None
            return row[idx]

        line_vals = []
        empty_streak = 0
        for offset, row in enumerate(
            ws.iter_rows(min_row=start_row, max_row=start_row + MAX_DATA_ROWS - 1, values_only=True)
        ):
            row_number = start_row + offset
            mapped = {key: cell(row, key) for key in cols}
            if all((v is None or str(v).strip() == "") for v in mapped.values()):
                empty_streak += 1
                if empty_streak >= 2:
                    break
                continue
            empty_streak = 0
            line_vals.append(self._parse_row(row_number, mapped))
        wb.close()

        if not line_vals:
            raise UserError(_("El archivo no contiene filas de datos."))
        self.write({"state": "validated"})
        self.env["step.management.crop.program.import.line"].sudo().create([
            dict(vals, import_id=self.id) for vals in line_vals
        ])
        self.invalidate_recordset(["line_ids", "line_count", "error_count", "ok_count"])
        self.message_post(body=_(
            "Validación: %(total)s fila(s), %(err)s con error."
        ) % {"total": self.line_count, "err": self.error_count})

    def _parse_row(self, row_number, mapped):
        Product = self.env["product.product"]
        Uom = self.env["uom.uom"]
        company = self.company_id

        raw = {
            "raw_product": mapped.get("product"), "raw_dose_per_ha": mapped.get("dose_per_ha"),
            "raw_uom": mapped.get("uom"), "raw_target": mapped.get("target"),
            "raw_week_number": mapped.get("week_number"), "raw_phi_days": mapped.get("phi_days"),
            "raw_rei_hours": mapped.get("rei_hours"),
        }
        vals = {k: ("" if v is None else str(v)) for k, v in raw.items()}
        vals["row_number"] = row_number

        errors = []
        for k, v in mapped.items():
            if isinstance(v, str) and v.strip().startswith("="):
                errors.append(_("La celda «%s» contiene una fórmula.") % k)

        raw_product = (str(mapped.get("product")).strip() if mapped.get("product") else "")
        if not raw_product:
            errors.append(_("Falta el producto."))
        else:
            products = Product.search([
                ("company_id", "in", [False, company.id]),
                "|", ("default_code", "=", raw_product), ("name", "=", raw_product),
            ], limit=2)
            if not products:
                errors.append(_("Producto «%s» no encontrado.") % raw_product)
            elif len(products) > 1:
                errors.append(_("Producto «%s» ambiguo.") % raw_product)
            else:
                vals["product_id"] = products.id

        dose = _to_float(mapped.get("dose_per_ha"))
        if dose is None or dose < 0:
            errors.append(_("Dosis por hectárea inválida."))
        else:
            vals["dose_per_ha"] = dose

        product = Product.browse(vals["product_id"]) if vals.get("product_id") else Product.browse()
        raw_uom = (str(mapped.get("uom")).strip() if mapped.get("uom") else "")
        if raw_uom:
            uoms = Uom.search([("name", "=", raw_uom)], limit=2)
            if not uoms:
                errors.append(_("Unidad de medida «%s» no encontrada.") % raw_uom)
            elif len(uoms) > 1:
                errors.append(_("Unidad de medida «%s» ambigua.") % raw_uom)
            elif product and uoms.category_id != product.uom_id.category_id:
                # R2: la UdM informada debe ser de la misma categoría que la
                # del producto; si no, no hay forma de consolidarla luego.
                errors.append(_(
                    "La unidad de medida «%(uom)s» no es de la misma "
                    "categoría que la del producto (%(target)s)."
                ) % {"uom": uoms.display_name, "target": product.uom_id.display_name})
            else:
                vals["uom_id"] = uoms.id
        elif product:
            # R2: sin UdM informada, se asume la UdM del producto.
            vals["uom_id"] = product.uom_id.id

        week, week_err = _to_int_strict(mapped.get("week_number"))
        if week_err:
            errors.append(_(
                "La semana debe ser un número entero (sin decimales)."
            ))
        elif week is not None:
            if not 1 <= week <= 53:
                errors.append(_("La semana debe estar entre 1 y 53."))
            else:
                vals["week_number"] = week

        phi, phi_err = _to_int_strict(mapped.get("phi_days"))
        if phi_err:
            errors.append(_(
                "La carencia debe ser un número entero de días (sin "
                "decimales)."
            ))
        elif phi is not None:
            if phi < 0:
                errors.append(_("La carencia no puede ser negativa."))
            else:
                vals["phi_days"] = phi

        rei, rei_err = _to_int_strict(mapped.get("rei_hours"))
        if rei_err:
            errors.append(_(
                "El reingreso debe ser un número entero de horas (sin "
                "decimales)."
            ))
        elif rei is not None:
            if rei < 0:
                errors.append(_("El reingreso no puede ser negativo."))
            else:
                vals["rei_hours"] = rei

        vals["target"] = (str(mapped.get("target")).strip() if mapped.get("target") else False)

        content = "|".join(_norm(mapped.get(k)) for k in sorted(mapped))
        vals["line_hash"] = hashlib.sha256(
            ("%s|%s|%s" % (self.file_hash or "", row_number, content)).encode("utf-8")
        ).hexdigest()

        vals["state"] = "error" if errors else "ok"
        vals["error"] = "\n".join(errors)
        return vals

    # ------------------------------------------------------------------
    # Importación (todo o nada) — crea el programa en borrador
    # ------------------------------------------------------------------
    def action_import(self):
        self.ensure_one()
        if self.state != "validated":
            raise UserError(_("Primero valide el archivo."))
        if self.error_count and not self.import_valid_only:
            raise UserError(_(
                "Hay %s fila(s) con error. Corríjalas y vuelva a validar, o marque "
                "«Importar sólo filas válidas»."
            ) % self.error_count)

        # R5: comprobación rápida sobre la intención completa (empresa,
        # temporada, tipo, centros, archivo) — no sólo empresa+archivo, para
        # no impedir reutilizar una receta legítima en otra temporada o con
        # otro conjunto de centros. El resguardo real ante concurrencia es
        # el índice único parcial (`init()`), no este `search()`.
        duplicate = self.search([
            ("id", "!=", self.id), ("company_id", "=", self.company_id.id),
            ("import_key_hash", "=", self.import_key_hash), ("state", "=", "imported"),
        ], limit=1)
        if duplicate:
            raise UserError(_(
                "Ya existe una importación con la misma empresa, temporada, "
                "tipo de programa, centros y archivo: %s. No se vuelve a "
                "cargar."
            ) % duplicate.name)

        rows = self.line_ids.filtered(lambda l: l.state == "ok")
        if not rows:
            raise UserError(_("No hay filas válidas para importar."))

        line_cmds = [
            (0, 0, {
                "sequence": (idx + 1) * 10,
                "product_id": row.product_id.id,
                "dose_per_ha": row.dose_per_ha,
                "uom_id": row.uom_id.id or False,
                "target": row.target or False,
                "week_number": row.week_number or 0,
                "phi_days": row.phi_days or 0,
                "rei_hours": row.rei_hours or 0,
            })
            for idx, row in enumerate(rows)
        ]

        # R4: si el encabezado no informó variedad, se infiere de los
        # centros cuando es inequívoca (nunca sobrescribe un valor ya
        # informado). `check_center_variety_coherence` ya validó arriba
        # (constrains) que la combinación es coherente.
        variety = self.variety or check_center_variety_coherence(
            self.center_ids, self.variety
        )

        # «Aprobación segura»: el programa importado siempre nace en
        # borrador; el flujo normal de cálculo/aprobación sigue vigente.
        # R5: la creación del programa y la marca "imported" van en el mismo
        # savepoint que el índice único parcial protege — si dos
        # importaciones con la misma intención corren en paralelo, la
        # segunda revierte aquí (no sólo pierde la carrera del `search()`
        # de arriba, que es de sólo lectura y no bloquea).
        try:
            with self.env.cr.savepoint():
                program = self.env["step.management.crop.program"].create({
                    "company_id": self.company_id.id,
                    "program_type": self.program_type,
                    "season": self.season,
                    "species": self.species,
                    "variety": variety,
                    "date": self.date,
                    "currency_id": self.currency_id.id,
                    "price_policy": self.price_policy,
                    "center_ids": [(6, 0, self.center_ids.ids)],
                    "line_ids": line_cmds,
                    "notes": _("Creado por importación %s.") % self.name,
                })
                self.write({"state": "imported", "program_id": program.id})
        except IntegrityError as exc:
            raise UserError(_(
                "Ya existe una importación con la misma empresa, temporada, "
                "tipo de programa, centros y archivo. No se vuelve a cargar."
            )) from exc
        rows.sudo().write({"imported": True})
        self.message_post(body=_(
            "Importado a %(program)s: %(lines)s línea(s) de receta, "
            "%(rows)s fila(s) del Excel."
        ) % {"program": program.name, "lines": len(line_cmds), "rows": len(rows)})
        program.message_post(body=_("Creado por importación %s.") % self.name)
        return {
            "type": "ir.actions.act_window",
            "res_model": "step.management.crop.program",
            "res_id": program.id, "view_mode": "form", "target": "current",
        }


class StepManagementCropProgramImportLine(models.Model):
    _name = "step.management.crop.program.import.line"
    _description = "Fila de carga de receta de programa"
    _order = "row_number, id"
    _check_company_auto = True

    import_id = fields.Many2one(
        "step.management.crop.program.import", required=True, ondelete="cascade", index=True,
    )
    company_id = fields.Many2one(related="import_id.company_id", store=True, index=True)
    row_number = fields.Integer(string="Fila")
    state = fields.Selection([("ok", "OK"), ("error", "Error")], default="ok", index=True)
    error = fields.Text(string="Detalle del error")
    imported = fields.Boolean(string="Importada", default=False)
    line_hash = fields.Char(size=64, index=True)

    raw_product = fields.Char(string="Producto (texto)")
    raw_dose_per_ha = fields.Char(string="Dosis/Ha (texto)")
    raw_uom = fields.Char(string="UdM (texto)")
    raw_target = fields.Char(string="Objetivo (texto)")
    raw_week_number = fields.Char(string="Semana (texto)")
    raw_phi_days = fields.Char(string="Carencia (texto)")
    raw_rei_hours = fields.Char(string="Reingreso (texto)")

    product_id = fields.Many2one("product.product", check_company=True)
    dose_per_ha = fields.Float(digits=(16, 4))
    uom_id = fields.Many2one("uom.uom")
    target = fields.Char()
    week_number = fields.Integer()
    phi_days = fields.Integer()
    rei_hours = fields.Integer()
