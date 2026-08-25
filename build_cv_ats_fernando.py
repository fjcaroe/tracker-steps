from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


OUTPUT = Path(r"C:\Users\tito4\Documents\Buscador trabajo Fernando\cv\CV ATS Fernando Caro Escobar.docx")


NAVY = RGBColor(31, 56, 85)
CHARCOAL = RGBColor(40, 40, 40)
MID = RGBColor(90, 90, 90)


def set_cell_free_page(section):
    section.top_margin = Cm(1.45)
    section.bottom_margin = Cm(1.35)
    section.left_margin = Cm(1.65)
    section.right_margin = Cm(1.65)
    section.header_distance = Cm(0.5)
    section.footer_distance = Cm(0.55)


def set_repeat_table_header(_):
    # Deliberately unused: ATS document contains no tables.
    return


def shade_paragraph(paragraph, fill="E8EEF4"):
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)


def add_bottom_border(paragraph, color="1F3855", size="8", space="2"):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), space)
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)


def set_keep_with_next(paragraph, keep=True):
    p_pr = paragraph._p.get_or_add_pPr()
    existing = p_pr.find(qn("w:keepNext"))
    if keep and existing is None:
        p_pr.append(OxmlElement("w:keepNext"))
    elif not keep and existing is not None:
        p_pr.remove(existing)


def set_keep_lines(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    if p_pr.find(qn("w:keepLines")) is None:
        p_pr.append(OxmlElement("w:keepLines"))


def add_section_heading(doc, text):
    p = doc.add_paragraph()
    p.style = doc.styles["Heading 1"]
    p.paragraph_format.space_before = Pt(7)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text.upper())
    r.bold = True
    r.font.name = "Arial"
    r.font.size = Pt(10.5)
    r.font.color.rgb = NAVY
    add_bottom_border(p, size="6", space="1")
    return p


def add_role(doc, title, company, dates, location="Santiago, Chile"):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3.5)
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(title)
    r.bold = True
    r.font.size = Pt(10)
    r.font.color.rgb = CHARCOAL
    r = p.add_run(f" | {company}")
    r.bold = True
    r.font.size = Pt(10)
    r.font.color.rgb = NAVY
    r = p.add_run(f" | {dates} | {location}")
    r.font.size = Pt(9.1)
    r.font.color.rgb = MID
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Cm(0.45)
    p.paragraph_format.first_line_indent = Cm(-0.28)
    p.paragraph_format.space_after = Pt(1.4)
    p.paragraph_format.line_spacing = 1.0
    set_keep_lines(p)
    r = p.add_run(text)
    r.font.name = "Arial"
    r.font.size = Pt(9.1)
    r.font.color.rgb = CHARCOAL
    return p


def add_compact_line(doc, label, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.0
    r = p.add_run(label)
    r.bold = True
    r.font.name = "Arial"
    r.font.size = Pt(9.2)
    r.font.color.rgb = NAVY
    r = p.add_run(text)
    r.font.name = "Arial"
    r.font.size = Pt(9.2)
    r.font.color.rgb = CHARCOAL
    return p


doc = Document()
for section in doc.sections:
    set_cell_free_page(section)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Arial"
normal.font.size = Pt(9.2)
normal.font.color.rgb = CHARCOAL
normal.paragraph_format.space_after = Pt(2)
normal.paragraph_format.line_spacing = 1.0

for style_name in ("List Bullet", "List Paragraph"):
    styles[style_name].font.name = "Arial"
    styles[style_name].font.size = Pt(9.1)

# ATS-safe masthead: plain paragraphs, no header object, no table, no icons.
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(1)
r = p.add_run("FERNANDO CARO ESCOBAR")
r.bold = True
r.font.name = "Arial"
r.font.size = Pt(20)
r.font.color.rgb = NAVY

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(2)
r = p.add_run("TECHNICAL LEAD | FULL STACK SOFTWARE ENGINEER | ARQUITECTURA E INTEGRACIONES")
r.bold = True
r.font.name = "Arial"
r.font.size = Pt(10.2)
r.font.color.rgb = CHARCOAL

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(6)
r = p.add_run("Santiago, Chile | +56 9 4410 0177 | fernandocaro1198@gmail.com | linkedin.com/in/fernandocaroe")
r.font.name = "Arial"
r.font.size = Pt(9)
r.font.color.rgb = MID
add_bottom_border(p, size="8", space="3")

add_section_heading(doc, "Perfil profesional")
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(3)
p.paragraph_format.line_spacing = 1.04
r = p.add_run(
    "Ingeniero en Informática y Software Engineer con más de 10 años de experiencia en desarrollo de software, "
    "integraciones y Business Intelligence para retail y banca. Más de 6 años en Falabella construyendo y "
    "evolucionando microservicios, servicios REST y SOA, aplicaciones Full Stack e integraciones corporativas. "
    "Experiencia en análisis técnico, migración de sistemas legados, CI/CD, datos y soporte a plataformas en "
    "producción. Reconocido consistentemente por pares y jefatura por su solidez técnica, proactividad, colaboración, "
    "cumplimiento y disposición para compartir conocimiento."
)
r.font.name = "Arial"
r.font.size = Pt(9.35)
r.font.color.rgb = CHARCOAL

add_section_heading(doc, "Competencias clave")
add_compact_line(doc, "Arquitectura y desarrollo: ", "Microservicios, arquitectura SOA, API REST, integración de sistemas, migración de legados, Full Stack, patrones de diseño, análisis técnico y resolución de problemas.")
add_compact_line(doc, "Liderazgo técnico: ", "Acompañamiento técnico, colaboración multidisciplinaria, revisión de soluciones, transferencia de conocimiento, mejora continua y orientación a calidad y entrega.")
add_compact_line(doc, "Tecnologías: ", "Node.js, JavaScript, React, Python, Java, .NET, Express, Koa, Oracle Service Bus, WebLogic, Oracle Data Integrator, Pentaho, QlikView.")
add_compact_line(doc, "Cloud, DevOps y datos: ", "GCP, Azure, Docker, Kubernetes, GitLab, GitHub, CI/CD, Oracle, SQL Server, MongoDB, Redis, Cosmos DB, Elasticsearch/Kibana.")

add_section_heading(doc, "Experiencia profesional")
add_role(doc, "Software Engineer | Full Stack e Integraciones", "Falabella", "mar 2020 – actualidad")
add_bullet(doc, "Analiza, prepara e implementa microservicios, soluciones SOA y servicios REST para sistemas corporativos de Recursos Humanos y experiencia de colaboradores.")
add_bullet(doc, "Desarrolla backend con Node.js y frameworks Express/Koa, junto con soluciones Full Stack en React, integrando servicios y fuentes de datos empresariales.")
add_bullet(doc, "Participa en migraciones masivas con Pentaho y en la modernización de aplicaciones legadas Oracle 12c hacia arquitecturas basadas en servicios y microservicios.")
add_bullet(doc, "Implementa y mantiene integraciones con Oracle Service Bus, WebLogic y Oracle Data Integrator, además de flujos CI/CD y despliegues sobre entornos cloud y contenedores.")
add_bullet(doc, "Actúa como referente técnico colaborativo: apoya a sus pares, comparte conocimiento y facilita entregas de calidad dentro de equipos multidisciplinarios.")

add_role(doc, "Desarrollador Node.js y Consultor de Integraciones", "ITAUM Consultorías y Soluciones Tecnológicas", "jun 2018 – feb 2020")
add_bullet(doc, "Desarrolló backend Node.js para Click & Collect corporativo, incluyendo BFF e integraciones con componentes de Warehouse Management.")
add_bullet(doc, "Construyó soluciones Full Stack y servicios REST con Node.js, Express, React, MongoDB, Redis, Oracle, Azure SQL y Cosmos DB.")
add_bullet(doc, "Integró servicios Azure Service Bus, Topics, Subscriptions y Blob Storage; trabajó con Elasticsearch/Kibana, Docker y Kubernetes.")
add_bullet(doc, "Prestó consultoría Business Intelligence con Pentaho Enterprise para Falabella, apoyando procesos de datos, soporte y operación de plataforma.")

# Start the second page at a complete role boundary for stable ATS parsing and readability.
p = doc.add_paragraph()
p.add_run().add_break(WD_BREAK.PAGE)
add_section_heading(doc, "Experiencia profesional (continuación)")

add_role(doc, "Desarrollador Business Intelligence", "Kibernum | BancoEstado", "mar 2016 – jun 2018")
add_bullet(doc, "Desarrolló y dio soporte a soluciones BI y procesos ETL con Pentaho, SQL Server, SSIS, QlikView y herramientas de administración de plataformas.")
add_bullet(doc, "Aplicó prácticas BI Agile e ITIL en procesos vinculados a arquitectura tecnológica, reporting y continuidad operacional.")
add_bullet(doc, "Impartió capacitación técnica en SQL, .NET, SSIS, procesos batch, Team Foundation Server, FAS y entornos mainframe.")

add_role(doc, "Analista QlikView BI", "Kibernum | BancoEstado", "oct 2015 – feb 2016")
add_bullet(doc, "Analizó y desarrolló soluciones QlikView para áreas de operaciones, sistemas y normalización de créditos.")
add_bullet(doc, "Implementó consultas T-SQL, cubos, DTS y flujos ETL desde Data Warehouse a Data Mart mediante SSIS, bcp y cmdsql.")

add_role(doc, "Socio | Consultor y Relator BI", "Karo OTEC", "may 2014 – sep 2015")
add_bullet(doc, "Realizó consultoría y capacitación en QlikView, bases de datos, herramientas de productividad y fundamentos de Business Intelligence.")

add_section_heading(doc, "Educación")
add_compact_line(doc, "Ingeniería en Informática | ", "INACAP, 2011–2014. Ayudante de Cálculo y mentor en actividades estudiantiles de tecnología.")
add_compact_line(doc, "Estudios de Ingeniería Civil en Computación e Informática | ", "Universidad Mayor, 2009–2010.")
add_compact_line(doc, "Intérprete Superior Musical en Piano | ", "Universidad Mayor, 2004–2012.")

add_section_heading(doc, "Certificaciones y formación relevante")
add_compact_line(doc, "Certificaciones: ", "Scrum Foundation Professional Certificate (CertiProf); Seguridad NCh-ISO 27001 para Profesionales TI; ITIL para Administración de Proyectos TI; Pentaho Enterprise Edition; QlikView Developer.")
add_compact_line(doc, "Formación técnica: ", "Python Machine Learning, Scikit-learn y TensorFlow; análisis y visualización de datos en R; Docker; Node.js; React; Ingeniería de Datos con Python; SOA; Ruby on Rails; Salesforce.")

# Simple ATS-safe footer with page numbering.
for section in doc.sections:
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Fernando Caro Escobar | CV ATS | ")
    r.font.name = "Arial"
    r.font.size = Pt(8)
    r.font.color.rgb = MID
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    r._r.append(fld_char1)
    r._r.append(instr_text)
    r._r.append(fld_char2)

doc.core_properties.title = "CV ATS - Fernando Caro Escobar"
doc.core_properties.subject = "Technical Lead | Full Stack Software Engineer | Arquitectura e Integraciones"
doc.core_properties.author = "Fernando Caro Escobar"
doc.core_properties.keywords = "Technical Lead, Líder Técnico, Software Engineer, Full Stack, Arquitectura, Integraciones, Node.js, React, Microservicios, REST, SOA, Chile"

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
print(OUTPUT)
