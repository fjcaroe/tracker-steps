from pathlib import Path
import os
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
ARTIFACTS = Path(os.environ.get("STEPS_MANUAL_ARTIFACTS", str(Path.home() / ".codex/local-artifacts/tracker-manual")))
SHOTS = ARTIFACTS / "screens"
OUTPUT = ARTIFACTS / "Manual_Steps_Tracker.docx"

GREEN = "103F35"
GREEN_2 = "2F9E72"
LIME = "D8FF62"
RED = "D94835"
BLUE = "2E74B5"
LIGHT = "E8EEF5"
INK = "183029"
MUTED = "5B6E68"
WHITE = "FFFFFF"


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for key, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{key}"))
        if node is None:
            node = OxmlElement(f"w:{key}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def keep_with_next(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    node = OxmlElement("w:keepNext")
    p_pr.append(node)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    tr_pr.append(OxmlElement("w:cantSplit"))


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Página ")
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor.from_string(MUTED)
    fld_char = OxmlElement("w:fldChar")
    fld_char.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char, instr, fld_sep, text, fld_end])


def set_alt_text(inline_shape, title, description):
    doc_pr = inline_shape._inline.docPr
    doc_pr.set("title", title)
    doc_pr.set("descr", description)


def configure_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for name, size, color, before, after in (
        ("Title", 28, GREEN, 0, 12),
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 12, "1F4D78", 10, 5),
    ):
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for list_name in ("List Bullet", "List Number"):
        style = styles[list_name]
        style.font.name = "Calibri"
        style.font.size = Pt(11)
        style.paragraph_format.left_indent = Inches(0.375)
        style.paragraph_format.first_line_indent = Inches(-0.188)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.25

    if "Eyebrow" not in styles:
        eyebrow = styles.add_style("Eyebrow", WD_STYLE_TYPE.PARAGRAPH)
    else:
        eyebrow = styles["Eyebrow"]
    eyebrow.font.name = "Calibri"
    eyebrow.font.size = Pt(9)
    eyebrow.font.bold = True
    eyebrow.font.color.rgb = RGBColor.from_string(RED)
    eyebrow.paragraph_format.space_after = Pt(5)
    eyebrow.paragraph_format.keep_with_next = True

    if "Caption Manual" not in styles:
        caption = styles.add_style("Caption Manual", WD_STYLE_TYPE.PARAGRAPH)
    else:
        caption = styles["Caption Manual"]
    caption.font.name = "Calibri"
    caption.font.size = Pt(9)
    caption.font.italic = True
    caption.font.color.rgb = RGBColor.from_string(MUTED)
    caption.paragraph_format.space_before = Pt(4)
    caption.paragraph_format.space_after = Pt(8)

    if "Callout" not in styles:
        callout = styles.add_style("Callout", WD_STYLE_TYPE.PARAGRAPH)
    else:
        callout = styles["Callout"]
    callout.font.name = "Calibri"
    callout.font.size = Pt(11)
    callout.font.bold = True
    callout.font.color.rgb = RGBColor.from_string(GREEN)
    callout.paragraph_format.space_after = Pt(5)
    callout.paragraph_format.keep_with_next = True

    header = section.header
    p = header.paragraphs[0]
    p.text = "STEPS TRACKER  /  MANUAL DE USO Y PRESENTACIÓN"
    p.style = styles["Eyebrow"]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_page_number(section.footer.paragraphs[0])


def add_eyebrow(doc, text):
    return doc.add_paragraph(text.upper(), style="Eyebrow")


def add_intro(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(10)
    run = p.add_run(text)
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor.from_string(MUTED)
    return p


def add_bullets(doc, items):
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def add_steps(doc, items):
    for index, item in enumerate(items, start=1):
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.left_indent = Inches(0.375)
        paragraph.paragraph_format.first_line_indent = Inches(-0.188)
        paragraph.paragraph_format.space_after = Pt(4)
        paragraph.paragraph_format.line_spacing = 1.25
        paragraph.add_run(f"{index}. ").bold = True
        paragraph.add_run(item)


def add_callout(doc, title, body, color=LIME):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(6.5)
    cell = table.cell(0, 0)
    cell.width = Inches(6.5)
    cell_margins(cell, 120, 160, 120, 160)
    shade(cell, color)
    p = cell.paragraphs[0]
    p.style = "Callout"
    p.add_run(title)
    p2 = cell.add_paragraph(body)
    p2.paragraph_format.space_after = Pt(0)
    prevent_row_split(table.rows[0])
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_image(doc, filename, caption, alt, width=6.5):
    path = SHOTS / filename
    if not path.exists():
        raise FileNotFoundError(path)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_together = True
    shape = p.add_run().add_picture(str(path), width=Inches(width))
    set_alt_text(shape, caption, alt)
    cap = doc.add_paragraph(caption, style="Caption Manual")
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return shape


def page(doc, eyebrow, title, intro=None):
    if len(doc.paragraphs) > 1:
        doc.add_page_break()
    add_eyebrow(doc, eyebrow)
    doc.add_heading(title, level=1)
    if intro:
        add_intro(doc, intro)


def add_two_col_table(doc, rows, widths=(2.05, 4.45), header=None):
    total_rows = len(rows) + (1 if header else 0)
    table = doc.add_table(rows=total_rows, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.style = "Table Grid"
    table.columns[0].width = Inches(widths[0])
    table.columns[1].width = Inches(widths[1])
    offset = 0
    if header:
        for col, text in enumerate(header):
            cell = table.cell(0, col)
            cell.text = text
            shade(cell, LIGHT)
            for run in cell.paragraphs[0].runs:
                run.bold = True
                run.font.color.rgb = RGBColor.from_string(GREEN)
        offset = 1
    for index, row_data in enumerate(rows, start=offset):
        for col, text in enumerate(row_data):
            cell = table.cell(index, col)
            cell.width = Inches(widths[col])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell_margins(cell)
            cell.text = text
            if col == 0:
                cell.paragraphs[0].runs[0].bold = True
        prevent_row_split(table.rows[index])
    if header:
        prevent_row_split(table.rows[0])
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


doc = Document()
configure_document(doc)

# Editorial cover.
cover = doc.add_table(rows=2, cols=1)
cover.alignment = WD_TABLE_ALIGNMENT.CENTER
cover.autofit = False
cover.columns[0].width = Inches(6.5)
shade(cover.cell(0, 0), GREEN)
shade(cover.cell(1, 0), LIME)
for cell in (cover.cell(0, 0), cover.cell(1, 0)):
    cell.width = Inches(6.5)
    cell_margins(cell, 220, 260, 220, 260)
for row in cover.rows:
    prevent_row_split(row)
p = cover.cell(0, 0).paragraphs[0]
r = p.add_run("STEPS TRACKER")
r.bold = True; r.font.size = Pt(13); r.font.color.rgb = RGBColor.from_string(LIME)
p2 = cover.cell(0, 0).add_paragraph("Operaciones agrícolas y forestales bajo control")
p2.runs[0].font.color.rgb = RGBColor.from_string(WHITE)
p2.runs[0].font.size = Pt(11)
p = cover.cell(1, 0).paragraphs[0]
r = p.add_run("Manual de uso y presentación")
r.bold = True; r.font.size = Pt(28); r.font.color.rgb = RGBColor.from_string(GREEN)
p2 = cover.cell(1, 0).add_paragraph("Una guía clara para comprender la solución, usarla con confianza y explicar su valor a un posible cliente.")
p2.runs[0].font.size = Pt(14); p2.runs[0].font.color.rgb = RGBColor.from_string(INK)

doc.add_paragraph()
doc.add_paragraph("Versión preparada para demostración comercial", style="Eyebrow")
meta = doc.add_paragraph("20 de agosto de 2026  ·  Chile")
meta.runs[0].font.size = Pt(11)
meta.runs[0].font.color.rgb = RGBColor.from_string(MUTED)
doc.add_paragraph()
quote = doc.add_paragraph("“La plataforma transforma recorridos, partes de trabajo y datos de maquinaria en información que se puede revisar, comparar y usar para tomar decisiones.”")
quote.paragraph_format.left_indent = Inches(0.35)
quote.paragraph_format.right_indent = Inches(0.35)
quote.runs[0].italic = True
quote.runs[0].font.size = Pt(14)
quote.runs[0].font.color.rgb = RGBColor.from_string(GREEN)
add_callout(doc, "Para quién es este manual", "Está escrito en lenguaje sencillo. No hace falta conocer informática, GPS ni sistemas de gestión para seguirlo.", LIGHT)

page(doc, "Primero, la idea", "1. Qué es Steps Tracker", "Steps Tracker reúne en un solo lugar lo que ocurre en terreno: quién trabajó, con qué máquina, dónde, durante cuánto tiempo, cuánto avanzó y cuánto combustible utilizó.")
doc.add_heading("El problema que resuelve", level=2)
add_bullets(doc, [
    "Evita que la información quede dispersa entre papeles, planillas, mensajes y la memoria de las personas.",
    "Permite comprobar recorridos y superficies mediante mapas, sin perder la posibilidad de ingresar datos manualmente.",
    "Entrega indicadores y reportes para conversar con operaciones, administración y gerencia usando la misma información.",
    "Conserva un historial auditable y prepara los datos para integrarlos con Odoo.",
])
doc.add_heading("La promesa en una frase", level=2)
add_callout(doc, "Del trabajo en terreno a una decisión respaldada", "La plataforma captura, ordena, muestra y compara la operación; el usuario mantiene el control y puede revisar cada registro.")
doc.add_heading("Qué no pretende hacer", level=2)
add_bullets(doc, [
    "No reemplaza la experiencia del encargado de campo: la vuelve visible y comprobable.",
    "No obliga a confiar ciegamente en el GPS: existe ingreso manual y revisión histórica.",
    "No mezcla demostración con producción: el modo demo usa datos sintéticos y lo indica en pantalla.",
])

page(doc, "Orientación", "2. Cómo está organizada la plataforma", "El menú izquierdo acompaña todo el recorrido. Cada módulo responde una pregunta concreta.")
add_two_col_table(doc, [
    ("Resumen", "¿Qué está pasando hoy y qué requiere atención?"),
    ("Registro", "¿Cómo comienzo una labor real con GPS?"),
    ("Ingreso manual", "¿Cómo agrego o corrijo un trabajo revisado por una persona?"),
    ("Monitoreo", "¿Dónde está la flota y qué está haciendo?"),
    ("Odómetro", "¿Cuánto recorrido y uso acumula cada equipo?"),
    ("Indicadores", "¿Cuáles son los números principales de la jornada?"),
    ("Analítica", "¿Qué tendencias, comparaciones y oportunidades aparecen?"),
    ("Maestros", "¿Cómo administro predios, polígonos, labores, máquinas y otros datos base?"),
    ("Historial", "¿Qué ocurrió en una sesión anterior y por dónde pasó?"),
], header=("Módulo", "Pregunta que responde"))
add_callout(doc, "Regla simple", "Para consultar, use Resumen, Monitoreo, Indicadores, Analítica o Historial. Para capturar información, use Registro o Ingreso manual. Para configurar, use Maestros.", LIGHT)

page(doc, "Uso diario", "3. Resumen de jornada", "Es la primera lectura para un encargado o jefe. Muestra el estado general sin exigir revisar pantalla por pantalla.")
add_image(doc, "resumen.jpg", "Pantalla Resumen: operación, alertas, ubicaciones y actividad reciente.", "Vista general de Steps Tracker con indicadores de jornada, mapa y máquinas.")
doc.add_heading("Qué conviene mirar primero", level=2)
add_steps(doc, [
    "Revise las tarjetas superiores: máquinas trabajando, superficie, distancia y consumo.",
    "Lea ‘Atención requerida’ para detectar combustible bajo u otra situación pendiente.",
    "Seleccione una ubicación para enfocar el mapa y la flota de ese fundo.",
    "Use la actividad reciente para comprender los últimos eventos.",
])

page(doc, "Captura automática", "4. Registro de una labor con GPS", "Este módulo se usa al iniciar una operación real. Asocia operador, máquina, predio y labor antes de comenzar a registrar el recorrido.")
add_image(doc, "registro.jpg", "Registro: selección de datos operacionales antes de iniciar la sesión.", "Formulario para iniciar una labor con seguimiento GPS.")
add_steps(doc, [
    "Confirme la ubicación o fundo.",
    "Seleccione máquina, operador, predio o polígono y labor.",
    "Revise los datos antes de iniciar.",
    "Comience el registro y mantenga la aplicación disponible durante el trabajo.",
    "Al finalizar, cierre la sesión para que quede completa en historial y reportes.",
])
add_callout(doc, "Importante", "Una sesión que queda abierta puede aparecer como activa y afectar el porcentaje de cierre. Conviene revisar este punto al terminar la jornada.", "FFF3DF")

page(doc, "Captura supervisada", "5. Ingreso manual de información", "Está pensado para personas que prefieren revisar cada dato en pantalla, para reconstruir registros antiguos o para completar una operación que no pudo capturarse automáticamente.")
add_image(doc, "ingreso_manual.jpg", "Ingreso manual: formulario guiado para crear un registro revisado.", "Pantalla de ingreso manual de labores y datos operacionales.")
doc.add_heading("Cuándo usarlo", level=2)
add_bullets(doc, [
    "Digitalizar partes de trabajo en papel.",
    "Incorporar registros históricos.",
    "Completar datos que faltaron por señal, batería o disponibilidad del equipo.",
    "Corregir información después de una revisión autorizada.",
])
add_callout(doc, "Control humano", "El propósito no es duplicar el Historial. Aquí se crea información; en Historial se consulta y reproduce lo que ya existe.")

page(doc, "Operación en vivo", "6. Monitoreo de flota", "Permite observar máquinas, ubicaciones y estados en un mapa. Es útil para coordinación, seguridad y seguimiento de la jornada.")
add_image(doc, "monitoreo.jpg", "Monitoreo: mapa operativo y estado de la flota seleccionada.", "Mapa con máquinas, predios y datos de operación en vivo.")
add_steps(doc, [
    "Elija la ubicación que desea revisar.",
    "Identifique las máquinas en el mapa o en la lista.",
    "Seleccione un equipo para centrarlo y ver su información.",
    "Revise alertas de combustible, detención o actividad cuando aparezcan.",
])
add_callout(doc, "Lectura correcta", "La posición sirve como apoyo operacional. Antes de concluir que existe un problema, confirme señal, hora de la última actualización y contexto del trabajo.", LIGHT)

page(doc, "Uso acumulado", "7. Odómetro y mantenimiento", "Ayuda a comparar distancia y uso acumulado por equipo. Es una base útil para planificar servicios, revisar desvíos y respaldar costos.")
add_image(doc, "odometro.jpg", "Odómetro: lecturas acumuladas y revisión por máquina.", "Panel de odómetros y uso acumulado de equipos.")
doc.add_heading("Ejemplos de decisiones", level=2)
add_bullets(doc, [
    "Programar una mantención por horas o kilómetros.",
    "Detectar una variación inesperada respecto del período anterior.",
    "Respaldar distribución de costos entre centros o faenas.",
])

page(doc, "Lectura gerencial", "8. Indicadores principales", "Resume los resultados más importantes y permite conversar sobre avance, productividad y consumo sin entrar todavía al detalle de cada sesión.")
add_image(doc, "indicadores.jpg", "Indicadores: números principales y comparaciones de la operación.", "Tarjetas de indicadores operacionales y gráficos comparativos.")
add_callout(doc, "Buena práctica", "Un indicador no es una sentencia. Si cambia mucho, abra Analítica o Historial y busque la causa: labor, terreno, máquina, operador, clima o calidad del registro.", LIGHT)

page(doc, "Centro de reportes", "9. Analítica: una visión para decidir", "Analítica transforma sesiones individuales en tendencias y comparaciones. Se puede filtrar por toda la empresa o por una ubicación, y por 7 días, 30 días o todo el histórico.")
add_image(doc, "analitica_resumen.jpg", "Analítica — Resumen: KPIs, alertas explicadas, evolución y ranking.", "Reporte ejecutivo con superficie, productividad, consumo, cierre, alertas y gráficos.")
add_two_col_table(doc, [
    ("Superficie", "Hectáreas cubiertas dentro del alcance y período."),
    ("Productividad", "Hectáreas realizadas por hora efectiva."),
    ("Consumo específico", "Litros utilizados por hectárea."),
    ("Cierre", "Porcentaje de sesiones correctamente terminadas."),
], header=("Indicador", "Cómo leerlo"))

page(doc, "Centro de reportes", "10. Analítica de productividad y combustible", "Estas dos vistas ayudan a explicar por qué una operación rindió más o consumió más, sin reducir el análisis a un único número.")
add_image(doc, "analitica_productividad.jpg", "Productividad: comparación de equipos y composición por labor.", "Gráficos de productividad, superficie por máquina, labor y predio.", width=5.55)
add_image(doc, "analitica_combustible.jpg", "Combustible: litros usados y consumo por hectárea.", "Gráficos de consumo total y específico por fecha y máquina.", width=5.55)
add_callout(doc, "Pregunta útil", "No pregunte solamente ‘¿quién consumió más?’. Compare litros por hectárea, tipo de labor, superficie, terreno y horas antes de tomar una decisión.", "FFF3DF")

page(doc, "Centro de reportes", "11. Analítica de flota y detalle auditable", "Uso de flota muestra carga productiva y disponibilidad. Detalle permite llegar hasta las sesiones que sostienen los gráficos y exportarlas a CSV.")
add_image(doc, "analitica_flota.jpg", "Uso de flota: horas productivas, capacidad disponible y estado.", "Gráficos de utilización, disponibilidad y combustible de la flota.", width=5.55)
add_image(doc, "analitica_detalle.jpg", "Detalle: sesiones filtradas y exportación CSV.", "Tabla de sesiones con fechas, máquinas, operadores, predios y resultados.", width=5.55)
add_callout(doc, "De gráfico a evidencia", "Use el reporte atractivo para descubrir una situación; use Detalle e Historial para comprobarla y explicarla.")

page(doc, "Configuración", "12. Maestros: la base de datos operacionales", "Los maestros evitan escribir nombres distintos para la misma cosa. Aquí se administran predios, polígonos, actividades, labores, implementos, máquinas y otros datos reutilizables.")
add_image(doc, "maestros.jpg", "Maestros: selección del mantenedor y lista de registros configurables.", "Panel de administración de datos maestros de Steps Tracker.")
add_steps(doc, [
    "Elija el maestro que desea administrar.",
    "Busque un registro existente antes de crear otro.",
    "Abra el registro para revisar o editar sus datos.",
    "Guarde y confirme que el cambio aparece en la lista.",
])
add_callout(doc, "Evite duplicados", "Antes de crear ‘Tractor 01’, ‘tractor 1’ o ‘Tractor N°1’, acuerde una forma única de nombrar los equipos.", LIGHT)

page(doc, "Configuración cartográfica", "13. Predios y polígonos", "Un polígono representa el contorno de un predio, lote o cuartel. Permite relacionar recorridos, superficies y reportes con un área concreta.")
add_image(doc, "editor_poligono.jpg", "Editor de polígonos: formulario y mapa en una distribución compacta.", "Mapa satelital con puntos editables que forman el contorno de un predio.")
add_steps(doc, [
    "Asigne un nombre claro y, si corresponde, un centro de costo.",
    "Haga clic en el mapa para agregar vértices alrededor del contorno.",
    "Arrastre un punto si necesita corregirlo.",
    "Use quitar punto o limpiar sólo cuando corresponda.",
    "Guarde y vuelva a abrir el registro para confirmar el contorno.",
])
add_callout(doc, "Consejo práctico", "Use pocos puntos pero bien ubicados. Un contorno simple es más fácil de revisar y mantener que uno con vértices innecesarios.", LIGHT)

page(doc, "Trazabilidad", "14. Historial y reproducción", "Historial sirve para consultar registros ya creados. Desde allí se puede abrir una sesión y reproducir su recorrido sobre el mapa.")
add_image(doc, "historial.jpg", "Historial: búsqueda, filtros y acceso al detalle de sesiones.", "Listado histórico de operaciones con filtros y acciones de revisión.")
add_steps(doc, [
    "Filtre por fecha, máquina, operador, predio o estado.",
    "Abra la sesión que desea comprobar.",
    "Revise duración, distancia, superficie, combustible y datos asociados.",
    "Reproduzca la ruta para entender por dónde se desplazó la máquina.",
])
add_callout(doc, "Diferencia clave", "Ingreso manual crea registros. Historial busca, revisa y reproduce registros que ya existen.")

page(doc, "Método recomendado", "15. Rutina diaria de trabajo", "Una rutina sencilla mejora la calidad de la información y reduce correcciones posteriores.")
add_two_col_table(doc, [
    ("Antes de salir", "Confirmar maestros, máquinas, operadores, labores y polígonos."),
    ("Al comenzar", "Iniciar Registro o dejar preparado el parte para Ingreso manual."),
    ("Durante la jornada", "Revisar Monitoreo y atender alertas con contexto."),
    ("Al finalizar", "Cerrar sesiones, completar datos faltantes y revisar estados."),
    ("Cierre diario", "Mirar Resumen e Indicadores; investigar excepciones."),
    ("Semanal o mensual", "Usar Analítica, exportar detalle y conversar sobre mejoras."),
], header=("Momento", "Acción"))
doc.add_heading("Responsabilidades sugeridas", level=2)
add_bullets(doc, [
    "Operador: registrar o informar la labor con datos correctos.",
    "Supervisor: revisar sesiones, alertas y correcciones.",
    "Administrador: mantener maestros y permisos.",
    "Jefatura: interpretar indicadores y decidir acciones.",
])

page(doc, "Presentación comercial", "16. Cómo vender la idea", "La mejor demostración cuenta una historia breve: problema, operación visible, evidencia y decisión.")
doc.add_heading("Guion de cinco minutos", level=2)
add_steps(doc, [
    "Problema: ‘Hoy la información está repartida y cuesta comprobar qué ocurrió’. ",
    "Captura: muestre Registro y explique que también existe Ingreso manual para no excluir a nadie.",
    "Visibilidad: abra Monitoreo y Resumen para enseñar la operación y las alertas.",
    "Decisión: entre a Analítica, cambie de reporte y explique productividad, combustible y uso de flota.",
    "Respaldo: cierre con Detalle, Historial, maestros e integración con Odoo.",
])
doc.add_heading("Beneficios que el cliente entiende", level=2)
add_bullets(doc, [
    "Menos tiempo reuniendo datos.",
    "Mayor confianza en lo informado.",
    "Mejor planificación de maquinaria y mantenciones.",
    "Conversaciones de productividad y combustible con evidencia.",
    "Trazabilidad para administración y control de costos.",
])
add_callout(doc, "Mensaje comercial", "No venda ‘una pantalla con mapas’. Venda control operacional, memoria histórica y mejores decisiones, con una adopción gradual que conserva la revisión humana.")

page(doc, "Preguntas frecuentes", "17. Dudas comunes y respuestas simples", "Estas respuestas ayudan durante una demostración o los primeros días de uso.")
add_two_col_table(doc, [
    ("¿Debo confiar siempre en el GPS?", "No. El GPS ayuda, pero los datos se pueden revisar y existe ingreso manual."),
    ("¿Se pueden cargar datos antiguos?", "Sí, mediante Ingreso manual y los endpoints previstos para la integración."),
    ("¿Qué diferencia hay entre Indicadores y Analítica?", "Indicadores resume; Analítica compara, explica tendencias y permite llegar al detalle."),
    ("¿Qué pasa si una sesión queda abierta?", "Aparecerá activa hasta cerrarse o corregirse; conviene revisarla al terminar la jornada."),
    ("¿Los polígonos se pueden editar?", "Sí. Se abren en el mantenedor, se corrigen en el mapa y se guardan nuevamente."),
    ("¿Funciona con Odoo?", "La solución contempla modelos, sincronización y endpoints compatibles; al desplegar se deben activar y validar en el servidor."),
    ("¿Los datos de la demostración son reales?", "No. El modo demo indica que usa datos sintéticos y no escribe en producción."),
], header=("Pregunta", "Respuesta"))
doc.add_heading("Glosario mínimo", level=2)
add_bullets(doc, [
    "Sesión: período de trabajo registrado desde inicio hasta cierre.",
    "Maestro: lista administrable que se reutiliza en formularios.",
    "Polígono: contorno geográfico de un predio, lote o cuartel.",
    "ha/h: hectáreas realizadas por hora efectiva.",
    "L/ha: litros de combustible usados por hectárea.",
    "CSV: archivo que se puede abrir en Excel para análisis adicional.",
    "API: conexión controlada que permite intercambiar datos con Odoo u otros sistemas.",
])

page(doc, "Cierre", "18. La idea central", "Steps Tracker hace visible la operación sin quitarle importancia a las personas que conocen el terreno.")
add_callout(doc, "Capturar · Revisar · Comprender · Decidir", "Ese es el recorrido completo de la plataforma: registrar el trabajo, comprobarlo, convertirlo en información y usarlo para mejorar.")
doc.add_heading("Antes de una demostración", level=2)
add_bullets(doc, [
    "Abra la aplicación con anticipación y confirme que el modo demo está activo.",
    "Elija una ubicación con máquinas y sesiones visibles.",
    "Prepare una historia concreta: productividad, combustible o trazabilidad.",
    "No intente mostrar todo. Recorra cinco pantallas y deje espacio para preguntas.",
    "Aclare qué está listo, qué usa datos demo y qué se activará al desplegar backend y Odoo.",
])
doc.add_paragraph()
p = doc.add_paragraph("Manual preparado para acompañar la demostración y la adopción de Steps Tracker.")
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.runs[0].bold = True
p.runs[0].font.size = Pt(13)
p.runs[0].font.color.rgb = RGBColor.from_string(GREEN)

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
print(OUTPUT)
