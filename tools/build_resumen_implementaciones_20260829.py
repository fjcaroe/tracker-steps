from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path(r"C:\Users\tito4\Documents\Odoo\deliverables\Resumen_implementaciones_Odoo_28-29_agosto_2026.docx")

# Preset: standard_business_brief.
# Named brand override: Steps Green replaces the preset heading blue and is
# reused consistently in headings, table headers and the lead callout.
GREEN = "0B6B5C"
DARK_GREEN = "114B3A"
LIGHT_GREEN = "E7F3EF"
PALE_GREEN = "F3F8F6"
GRAY = "5F6B66"
LIGHT_GRAY = "F2F4F3"
MID_GRAY = "D8DEDB"
INK = "1B2521"
AMBER = "A56A00"
RED = "9B1C1C"
WHITE = "FFFFFF"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa, indent_dxa=120):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            width = widths_dxa[idx]
            cell.width = Inches(width / 1440)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)


def set_run(run, size=None, bold=None, color=None, italic=None, font="Calibri"):
    run.font.name = font
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), font)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), font)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def add_text(doc, text, *, bold_label=None, after=6, color=INK):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.10
    if bold_label and text.startswith(bold_label):
        label, rest = text[: len(bold_label)], text[len(bold_label) :]
        set_run(p.add_run(label), bold=True, color=color)
        set_run(p.add_run(rest), color=color)
    else:
        set_run(p.add_run(text), color=color)
    return p


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.paragraph_format.left_indent = Inches(0.5 if level == 0 else 0.75)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.167
    set_run(p.add_run(text), color=INK)
    return p


def add_number(doc, text, style="List Number"):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.167
    set_run(p.add_run(text), color=INK)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    set_run(p.add_run(text), color=GREEN if level < 3 else DARK_GREEN, bold=True)
    return p


def add_callout(doc, label, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(12)
    p.paragraph_format.left_indent = Inches(0.18)
    p.paragraph_format.right_indent = Inches(0.12)
    p.paragraph_format.line_spacing = 1.10
    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), LIGHT_GREEN)
    p_pr.append(shd)
    borders = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "24")
    left.set(qn("w:space"), "8")
    left.set(qn("w:color"), GREEN)
    borders.append(left)
    p_pr.append(borders)
    set_run(p.add_run(label + " "), bold=True, color=DARK_GREEN)
    set_run(p.add_run(text), color=INK)


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_status_table(doc, rows):
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    widths = [1900, 4220, 1700, 1540]
    headers = ["Frente", "Resultado principal", "Ambientes", "Estado"]
    for idx, text in enumerate(headers):
        cell = table.rows[0].cells[idx]
        set_cell_shading(cell, DARK_GREEN)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        set_run(p.add_run(text), size=9.5, bold=True, color=WHITE)
    repeat_header(table.rows[0])
    for r_idx, row_data in enumerate(rows):
        cells = table.add_row().cells
        for idx, value in enumerate(row_data):
            if r_idx % 2:
                set_cell_shading(cells[idx], PALE_GREEN)
            cells[idx].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cells[idx].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            color = DARK_GREEN if idx == 3 and value == "Completado" else INK
            set_run(p.add_run(value), size=9, bold=(idx in (0, 3)), color=color)
    set_table_geometry(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_pending_table(doc, rows):
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    widths = [1150, 2250, 3860, 2100]
    headers = ["Prioridad", "Área", "Pendiente", "Siguiente acción"]
    for idx, text in enumerate(headers):
        cell = table.rows[0].cells[idx]
        set_cell_shading(cell, DARK_GREEN)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        set_run(p.add_run(text), size=9.3, bold=True, color=WHITE)
    repeat_header(table.rows[0])
    for r_idx, row_data in enumerate(rows):
        cells = table.add_row().cells
        for idx, value in enumerate(row_data):
            if r_idx % 2:
                set_cell_shading(cells[idx], LIGHT_GRAY)
            p = cells[idx].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            priority_color = RED if value == "Alta" else AMBER if value == "Media" else GRAY
            set_run(p.add_run(value), size=8.7, bold=(idx in (0, 1)),
                    color=priority_color if idx == 0 else INK)
    set_table_geometry(table, widths)


def add_footer(section):
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    set_run(p.add_run("Steps Consulting  |  Resumen de implementaciones  |  28-29 agosto 2026"),
            size=8.5, color=GRAY)


def build():
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    add_footer(section)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10
    for name, size, before, after in (
        ("Heading 1", 16, 16, 8),
        ("Heading 2", 13, 12, 6),
        ("Heading 3", 12, 8, 4),
    ):
        st = styles[name]
        st.font.name = "Calibri"
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor.from_string(GREEN if name != "Heading 3" else DARK_GREEN)
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.keep_with_next = True

    # Memo masthead.
    kicker = doc.add_paragraph()
    kicker.paragraph_format.space_before = Pt(16)
    kicker.paragraph_format.space_after = Pt(7)
    set_run(kicker.add_run("INFORME DE AVANCE"), size=10, bold=True, color=GREEN)

    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(5)
    set_run(title.add_run("Resumen de implementaciones Odoo"), size=26, bold=True, color=DARK_GREEN)
    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(16)
    set_run(subtitle.add_run("Trabajo ejecutado entre el 28 y 29 de agosto de 2026"),
            size=13.5, color=GRAY)

    meta = [
        ("Proyecto", "Steps - aplicaciones agrícolas, Nómina y Contabilidad"),
        ("Ambientes", "Desarrollo, Demo y Demo-SyS, según compatibilidad"),
        ("Objetivo", "Convertir prototipos en módulos Odoo 18 homologados, seguros y demostrables"),
        ("Estado", "Cinco frentes principales implementados y validados"),
    ]
    for label, value in meta:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        set_run(p.add_run(label + ": "), size=10.5, bold=True, color=DARK_GREEN)
        set_run(p.add_run(value), size=10.5, color=INK)

    add_callout(
        doc,
        "Resultado ejecutivo.",
        "Se consolidaron Previred, contabilidad bimoneda, Maquinaria/BPA, la limpieza de aplicaciones Studio y Tesorería. Los despliegues dirigidos quedaron respaldados y sin mezclar los datos de las bases.",
    )

    add_heading(doc, "1. Panorama general", 1)
    add_status_table(doc, [
        ("Previred", "Dataset versionado y corrección de campos 13, 93 y 105.", "Dev / Demo / SyS", "Completado"),
        ("Bimoneda", "Segunda moneda persistida en asientos, pagos y reportes.", "Dev / Demo / SyS", "Completado"),
        ("Maquinaria y BPA", "OT correlativa, catálogo, costeo, contabilidad e integración BPA.", "Dev / Demo", "Completado"),
        ("Apps y Studio", "Fusión de menús, conservación de funciones únicas y retiro de duplicados.", "Dev / Demo", "Completado"),
        ("Tesorería", "Flujo de caja a cinco semanas, seguridad, migración y exportación XLSX.", "Dev / Demo / SyS", "Completado"),
    ])

    add_heading(doc, "2. Implementaciones realizadas", 1)

    add_heading(doc, "2.1 Previred y Nómina", 2)
    add_text(doc, "Se consolidó un módulo común de Previred y adaptadores para los motores de Nómina presentes en los distintos ambientes, evitando parches directos sobre código compartido de terceros.")
    add_bullet(doc, "Campo 13 calculado desde días de asistencia efectiva, excluyendo ausencias y licencias.")
    add_bullet(doc, "Campo 93 derivado del horario: jornada completa o parcial; migración inicial para horarios de hasta 30 horas.")
    add_bullet(doc, "Campo 105 obtenido desde la cuenta analítica o centro de costo del contrato, limitado a 20 caracteres.")
    add_bullet(doc, "Correcciones aplicadas a los perfiles Previred v84 y v98.")
    add_bullet(doc, "Despliegue en LAB_TAREAS, STEPS_DEMO y STEPS_DEMO_SYS; validación funcional con liquidación real de Demo-SyS.")
    add_text(doc, "Validación:", bold_label="Validación:")
    add_bullet(doc, "21 casos del dataset ejecutados sin fallas ni errores.", level=1)

    add_heading(doc, "2.2 Contabilidad bimoneda", 2)
    add_text(doc, "La segunda moneda dejó de ser sólo una conversión visual del reporte: ahora forma parte de la transacción contable y queda registrada por apunte.")
    add_bullet(doc, "Moneda operacional configurable por empresa.")
    add_bullet(doc, "Fecha de tasa, dólar observado, tasa por apunte y débito/crédito/saldo operacional persistidos.")
    add_bullet(doc, "Apunte operacional de diferencia de cambio cuando tasas distintas por línea descuadran el libro auxiliar, sin alterar el balance en la moneda principal.")
    add_bullet(doc, "Pagos con visualización y conversión en ambos sentidos.")
    add_bullet(doc, "Balance General y Estado de Resultados con columnas operacionales cuando el handler del reporte es compatible.")
    add_bullet(doc, "USD configurado como moneda operacional en Desarrollo y Demo; módulo instalado sin configuración global en Demo-SyS por su estructura multiempresa.")
    add_text(doc, "Validación:", bold_label="Validación:")
    add_bullet(doc, "Cinco flujos funcionales probados, sin fallas ni errores.", level=1)

    add_heading(doc, "2.3 Maquinaria agrícola e integración BPA", 2)
    add_text(doc, "Se reformuló el módulo de Maquinaria para reemplazar el comportamiento heredado de Studio y acercarlo al estilo funcional de las aplicaciones agrícolas Steps.")
    add_bullet(doc, "Número de OT correlativo, único por empresa y generado por el servidor.")
    add_bullet(doc, "Nombre automático con número, fecha, orden de trabajo y folio BPA, conservando el nombre histórico en migraciones.")
    add_bullet(doc, "Catálogo canónico de ocho conceptos de maquinaria, homologado por significado y código.")
    add_bullet(doc, "Correcciones de costo por hora, consumos, provisiones, mantenciones y contabilización.")
    add_bullet(doc, "Diario y cuentas contables de Maquinaria homologados entre Desarrollo y Demo.")
    add_bullet(doc, "Vínculo nativo con órdenes BPA; migraciones idempotentes y archivo reversible de la vista Studio anterior.")
    add_bullet(doc, "Permisos, aislamiento multiempresa y conservación de 20 encabezados, 32 líneas y 27 conciliaciones existentes.")
    add_bullet(doc, "Logo moderno mantenido y eliminación completa de la imagen con iniciales JE.")
    add_text(doc, "Versiones y validación:", bold_label="Versiones y validación:")
    add_bullet(doc, "step_machinery 18.0.21.0.0 y step_bpa_irrigation 18.0.2.3.1 en Desarrollo y Demo.", level=1)
    add_bullet(doc, "50 pruebas en base desechable y flujo controlado con rollback, sin fallas ni errores.", level=1)

    add_heading(doc, "2.4 Consolidación de aplicaciones y desarrollos Studio", 2)
    add_text(doc, "Se auditó la convivencia entre prototipos Studio y módulos propios. El criterio fue conservar la información y las funciones únicas, moverlas al módulo definitivo y archivar únicamente lo redundante.")
    add_bullet(doc, "Fusión de BPA y Riego, QA-Inspecciones, Protección Laboral y Fletes.")
    add_bullet(doc, "Veinte menús útiles trasladados, entradas equivalentes archivadas y raíces Studio vacías ocultadas sin borrar datos.")
    add_bullet(doc, "Steps Tracker quedó con un solo acceso publicado por step_tracker_odoo.")
    add_bullet(doc, "Gestión y Costos borrador fue absorbido después de conservar sus maestros y accesos no replicados.")
    add_bullet(doc, "En Fletes se retiraron equivalentes exactos; informes, servicios, camiones y modalidad de frío se conservaron cuando no existía reemplazo nativo.")
    add_bullet(doc, "Resultado verificado: cero aplicaciones raíz duplicadas en Desarrollo y Demo.")

    add_heading(doc, "2.5 Tesorería en Contabilidad", 2)
    add_text(doc, "Se implementó Tesorería como un núcleo común y dos puentes opcionales: integración agrícola para Fundo e integración con pagos por lotes. Demo-SyS usa el mismo núcleo sin instalar dependencias que no posee.")
    add_bullet(doc, "Maestro multiempresa de conceptos de flujo, con múltiples cuentas contables por concepto.")
    add_bullet(doc, "Flujo de caja de cinco semanas con Vencido, W1-W5 y Otros.")
    add_bullet(doc, "Fuentes automáticas: clientes, proveedores, notas de venta, órdenes de compra y proformas; además de ingresos y pagos manuales.")
    add_bullet(doc, "Saldo inicial desde diarios bancarios y de efectivo, con corte anterior al inicio del horizonte.")
    add_bullet(doc, "Dataset canónico compartido por hojas, resumen, indicadores y exportación XLSX.")
    add_bullet(doc, "Conversión a moneda del flujo y moneda auxiliar, incluyendo documentos en una tercera moneda.")
    add_bullet(doc, "Cuatro perfiles: consulta, operación, aprobación y configuración, con reglas multiempresa.")
    add_bullet(doc, "Migración idempotente de 12 conceptos y un flujo Studio; menús Studio archivados de forma reversible.")
    add_bullet(doc, "Corrección posterior del bloqueo ORM: un flujo aprobado ya no puede alterarse por vista, RPC ni importación; sólo aprueba el usuario designado.")
    add_bullet(doc, "Totales de moneda auxiliar derivados del detalle real, evitando conversiones agregadas inconsistentes.")
    add_text(doc, "Estado final:", bold_label="Estado final:")
    add_bullet(doc, "step_account_treasury 18.0.1.0.1, con código idéntico en los tres ambientes.", level=1)
    add_bullet(doc, "39 métodos de prueba ejecutados; Odoo contabilizó 51 verificaciones, con 0 fallas y 0 errores.", level=1)
    add_bullet(doc, "Validación web de la lista de Flujos de Caja en Desarrollo, Demo y Demo-SyS; servicios activos y HTTP 200.", level=1)

    add_heading(doc, "3. Seguridad, respaldo y homologación", 1)
    add_number(doc, "Se comprobó que no existieran actualizaciones de módulos, respaldos o despliegues concurrentes antes de intervenir.")
    add_number(doc, "Se generaron dumps verificables de cada base y copias de los addons antes de las actualizaciones dirigidas.")
    add_number(doc, "Los despliegues se realizaron de forma secuencial, reiniciando únicamente el servicio correspondiente.")
    add_number(doc, "No se copiaron ni mezclaron bases de datos; Demo-SyS conservó sus datos y adaptaciones de SimpleDigital/SyS.")
    add_number(doc, "No se ejecutaron pagos, correos, documentos tributarios, archivos bancarios ni otras acciones externas reales.")
    add_number(doc, "No se realizó git push; el código permanece local y desplegado en los ambientes autorizados.")

    add_heading(doc, "4. Pendientes", 1)
    add_callout(doc, "Criterio.", "Los siguientes puntos no invalidan lo entregado. Corresponden a definiciones funcionales, deuda técnica previa o mejoras que requieren datos y especificaciones adicionales.")
    add_pending_table(doc, [
        ("Alta", "Previred", "Contrastar el TXT final con el nuevo archivo oficial que entregará el cliente y obtener aceptación funcional.", "Ejecutar matriz campo a campo y prueba con nómina representativa."),
        ("Alta", "BPA", "Evitar división por cero cuando Capacidad lts aplicadora sea 0 y revisar el consumo por hora obligatorio heredado de Studio.", "Definir regla funcional y migrar campos a modelo nativo."),
        ("Alta", "Infraestructura", "steps_api figura instalado en Desarrollo y Demo, pero su código no existe en los addons_path.", "Recuperar el addon o desinstalar limpiamente tras respaldo y auditoría."),
        ("Media", "Maquinaria", "El diario CDMaq mantiene una lista cerrada de cuentas permitidas.", "Decidir entre lista abierta o validación dinámica por conceptos."),
        ("Media", "Tesorería", "Falta definir y construir el archivo bancario de pagos por lotes.", "Solicitar formato por banco y diseñar adaptadores versionados."),
        ("Media", "Tesorería", "No existe aún dashboard inicial OWL ni exportación PDF.", "Diseñar tarjetas, gráfico acumulado y reporte PDF desde el dataset canónico."),
        ("Media", "Tesorería", "Falta prueba de carga con volumen real y validación visual con usuario restringido.", "Crear dataset de estrés y ejecutar pruebas por perfil."),
        ("Media", "Bimoneda", "Demo-SyS no tiene moneda operacional definida por compañía y los datos históricos no fueron convertidos.", "Configurar empresa por empresa y aprobar una migración histórica si corresponde."),
        ("Media", "Tracker", "Los modelos step.tracker.* aún existen tanto en step_hr como en step_tracker_odoo.", "Separar modelos y migrar referencias en una refactorización controlada."),
        ("Baja", "Fletes/Studio", "Persisten funciones exclusivas de Studio sin equivalente nativo, conservadas para no perder capacidad.", "Reimplementar informes y maestros faltantes antes de archivar el legado."),
        ("Baja", "Ambientes", "Maquinaria no está instalada en Demo-SyS y steps_qa conserva una versión anterior.", "Definir si esos ambientes deben recibir el stack agrícola y homologarlos."),
        ("Baja", "Repositorios", "Los releases aún no fueron publicados en Git por instrucción operativa.", "Consolidar commits y publicar sólo cuando se autorice."),
    ])

    add_heading(doc, "5. Próximos pasos recomendados", 1)
    add_number(doc, "Cerrar la validación Previred apenas llegue el archivo oficial del cliente.", "List Number 2")
    add_number(doc, "Resolver primero la deuda BPA que impide crear OT cuando la capacidad está en cero.", "List Number 2")
    add_number(doc, "Sanear steps_api antes de ampliar los despliegues, para eliminar una referencia instalada sin código.", "List Number 2")
    add_number(doc, "Completar la experiencia comercial de Tesorería con dashboard y PDF, sin construir archivos bancarios hasta contar con especificación aprobada.", "List Number 2")
    add_number(doc, "Preparar una ronda de homologación final que incluya perfiles restringidos, carga representativa y evidencia visual por ambiente.", "List Number 2")

    add_callout(doc, "Conclusión.", "El avance de estas dos jornadas dejó módulos más cercanos a producto: transacciones persistentes, migraciones idempotentes, permisos, trazabilidad, integración entre aplicaciones y un release común por ambiente. Los pendientes están identificados y pueden abordarse sin rehacer lo ya construido.")

    props = doc.core_properties
    props.title = "Resumen de implementaciones Odoo - 28 y 29 de agosto de 2026"
    props.subject = "Avance técnico y funcional de módulos Steps"
    props.author = "Steps Consulting"
    props.keywords = "Odoo, Steps, Previred, Maquinaria, BPA, Tesorería, Bimoneda"
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
