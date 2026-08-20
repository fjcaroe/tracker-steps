from pathlib import Path

from docx import Document


SOURCE = Path("docs/Manual_Steps_Tracker.docx")
OUTPUT = Path("docs/Manual_Steps_Tracker.docx")


PARAGRAPH_REPLACEMENTS = {
    "Es la primera lectura para un encargado o jefe. Muestra el estado general sin exigir revisar pantalla por pantalla.":
        "Es la primera lectura para un encargado o jefe. En datos reales muestra máquinas, sesiones abiertas, horas productivas, distancia GPS, combustible estimado y la calidad de los registros visibles para el usuario.",
    "1. Revise las tarjetas superiores: máquinas trabajando, superficie, distancia y consumo.":
        "1. Revise las tarjetas superiores: máquinas, sesiones abiertas, horas, distancia GPS y combustible estimado.",
    "2. Lea ‘Atención requerida’ para detectar combustible bajo u otra situación pendiente.":
        "2. Revise la actividad reciente. Si aparece “Pendiente cierre”, la sesión lleva más de 24 horas abierta y no se suma a horas ni combustible productivo.",
    "3. Seleccione una ubicación para enfocar el mapa y la flota de ese fundo.":
        "3. Observe el índice de integridad: combina sesiones cerradas, presencia de GPS y asignación de labor y centro de costo.",
    "4. Use la actividad reciente para comprender los últimos eventos.":
        "4. Use Demo para explicar un escenario completo y Datos reales para revisar la operación autorizada.",
    "Permite observar máquinas, ubicaciones y estados en un mapa. Es útil para coordinación, seguridad y seguimiento de la jornada.":
        "En Datos reales muestra los polígonos registrados y ubica solamente sesiones abiertas durante las últimas 24 horas que tengan posición GPS. Las sesiones antiguas quedan como “Cierre pendiente” y no se presentan como máquinas activas.",
    "Ayuda a comparar distancia y uso acumulado por equipo. Es una base útil para planificar servicios, revisar desvíos y respaldar costos.":
        "Compara por máquina el odómetro GPS, las horas registradas, el combustible estimado y la cantidad de puntos. Las sesiones abiertas por más de 24 horas se mantienen visibles, pero no inflan horas ni combustible.",
    "Resume los resultados más importantes y permite conversar sobre avance, productividad y consumo sin entrar todavía al detalle de cada sesión.":
        "Resume resultados reales y calidad de información. El índice de integridad muestra cierre de sesiones, cobertura GPS y clasificación por labor y centro de costo; además compara horas, distancia, consumo horario y sesiones abiertas.",
    "Analítica transforma sesiones individuales en tendencias y comparaciones. Se puede filtrar por toda la empresa o por una ubicación, y por 7 días, 30 días o todo el histórico.":
        "Analítica transforma sesiones individuales en tendencias y comparaciones. En Datos reales usa sesiones, GPS, máquinas, labores y centros de costo; permite revisar 7 días, 30 días o todo el histórico cargado y exportar el detalle a CSV.",
    "Los maestros evitan escribir nombres distintos para la misma cosa. Aquí se administran predios, polígonos, actividades, labores, implementos, máquinas y otros datos reutilizables.":
        "Los maestros evitan escribir nombres distintos para la misma cosa. Todos los usuarios autenticados pueden consultarlos, pero sólo un administrador puede crear, editar o eliminar predios, polígonos, actividades, labores, implementos, máquinas y demás catálogos.",
    "Historial sirve para consultar registros ya creados. Desde allí se puede abrir una sesión y reproducir su recorrido sobre el mapa.":
        "Historial sirve para consultar registros ya creados. Las sesiones reales pueden abrirse para cargar sus puntos GPS, reproducir el recorrido sobre el mapa y revisar velocidad y posición a lo largo del tiempo.",
    "Administrador: mantener maestros y permisos.":
        "Administrador: mantener maestros y permisos; los demás perfiles pueden consultar los catálogos y operar sólo sobre sus centros de costo autorizados.",
    "3. Visibilidad: abra Monitoreo y Resumen para enseñar la operación y las alertas.":
        "3. Visibilidad: abra Monitoreo y Resumen; cambie a Datos reales para mostrar polígonos, sesiones y alertas de cierre sin mezclar el laboratorio.",
    "4. Decisión: entre a Analítica, cambie de reporte y explique productividad, combustible y uso de flota.":
        "4. Decisión: entre a Indicadores y Analítica para explicar integridad, horas, distancia, combustible, tendencias y uso de flota.",
    "5. Respaldo: cierre con Detalle, Historial, maestros e integración con Odoo.":
        "5. Respaldo: cierre con Detalle, reproducción GPS, maestros protegidos e integración activa con Odoo.",
    "Aclare qué está listo, qué usa datos demo y qué se activará al desplegar backend y Odoo.":
        "Aclare cuándo está usando Demo y cuándo Datos reales. Confirme que la API y la aplicación de Odoo están disponibles antes de una demostración productiva.",
}


def replace_paragraph_text(paragraph, new_text: str) -> None:
    if paragraph.runs:
        paragraph.runs[0].text = new_text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(new_text)


def set_cell(cell, title: str, body: str) -> None:
    while len(cell.paragraphs) < 2:
        cell.add_paragraph()
    replace_paragraph_text(cell.paragraphs[0], title)
    replace_paragraph_text(cell.paragraphs[1], body)
    for paragraph in cell.paragraphs[2:]:
        replace_paragraph_text(paragraph, "")


doc = Document(SOURCE)

found = set()
for paragraph in doc.paragraphs:
    original = paragraph.text
    if original in PARAGRAPH_REPLACEMENTS:
        replace_paragraph_text(paragraph, PARAGRAPH_REPLACEMENTS[original])
        found.add(original)

missing = set(PARAGRAPH_REPLACEMENTS) - found
if missing:
    raise RuntimeError(f"No se encontraron {len(missing)} párrafos esperados: {sorted(missing)}")

# Callouts operacionales.
set_cell(doc.tables[5].cell(0, 0), "Importante", "Una sesión abierta por más de 24 horas se marca como pendiente de cierre. Sigue en el historial y conserva su GPS, pero no suma horas ni combustible productivo hasta ser revisada.")
set_cell(doc.tables[7].cell(0, 0), "Lectura correcta", "Una máquina aparece en el mapa real sólo si tiene una sesión abierta reciente y una posición GPS. Si no aparece, revise la hora del último punto, la señal y el estado de cierre.")
set_cell(doc.tables[8].cell(0, 0), "Buena práctica", "El índice de integridad no califica a una persona. Señala si falta cerrar sesiones, capturar GPS o asignar labor y centro de costo para que los reportes sean comparables.")

# Tabla de indicadores productivos reales.
indicator_rows = [
    ("Indicador", "Cómo leerlo"),
    ("Horas registradas", "Duración efectiva; excluye sesiones abiertas por más de 24 horas."),
    ("Distancia GPS", "Suma del recorrido almacenado por máquina o período."),
    ("Consumo horario", "Litros estimados por hora según configuración de máquina y labor."),
    ("Integridad", "Promedio de cierre, cobertura GPS y clasificación completa."),
]
for row, values in zip(doc.tables[9].rows, indicator_rows):
    replace_paragraph_text(row.cells[0].paragraphs[0], values[0])
    replace_paragraph_text(row.cells[1].paragraphs[0], values[1])

set_cell(doc.tables[10].cell(0, 0), "Pregunta útil", "No pregunte solamente “¿quién consumió más?”. Compare horas, distancia, consumo horario, labor, terreno y calidad del registro antes de decidir.")
set_cell(doc.tables[11].cell(0, 0), "De gráfico a evidencia", "Use Analítica para descubrir una situación; use Detalle, Historial y reproducción GPS para comprobarla. Los datos Demo y reales siempre se mantienen separados.")
set_cell(doc.tables[14].cell(0, 0), "Diferencia clave", "Ingreso manual crea partes revisados. Historial consulta y reproduce registros existentes. Maestros configura catálogos y sólo permite cambios a administradores.")

# Preguntas frecuentes: actualizar respuestas y agregar permisos.
faq = doc.tables[17]
for row in faq.rows[1:]:
    question = row.cells[0].text.strip()
    if question == "¿Qué pasa si una sesión queda abierta?":
        replace_paragraph_text(row.cells[1].paragraphs[0], "Si supera 24 horas se marca “Pendiente cierre” y deja de sumar horas y combustible productivo, pero no se borra ni pierde su GPS.")
    elif question == "¿Funciona con Odoo?":
        replace_paragraph_text(row.cells[1].paragraphs[0], "Sí. La aplicación step_hr está activa en Odoo y sincroniza máquinas, predios, sesiones y partes mediante la API de Tracker.")

new_row = faq.add_row()
replace_paragraph_text(new_row.cells[0].paragraphs[0], "¿Quién puede modificar los maestros?")
replace_paragraph_text(new_row.cells[1].paragraphs[0], "Sólo un administrador. Los usuarios autenticados pueden consultar catálogos y trabajar dentro de sus centros de costo autorizados.")

doc.core_properties.title = "Manual de uso y presentación comercial - Steps Tracker"
doc.core_properties.subject = "Guía práctica de operación, analítica, maestros, historial e integración con Odoo"
doc.core_properties.comments = "Actualizado con paneles productivos reales, seguridad y reproducción GPS."
doc.save(OUTPUT)
print(f"Manual actualizado: {OUTPUT}")
