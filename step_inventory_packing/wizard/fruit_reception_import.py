"""Importación controlada de tarjas recibidas desde una tabla Excel plana."""

import base64
from io import BytesIO

from openpyxl import load_workbook

from odoo import fields, models, _
from odoo.exceptions import UserError


_HEADERS = ("Tarja", "Producto", "Cantidad", "Kilos", "Cajas", "Peso bruto", "Destare")


class StepFruitReceptionImport(models.TransientModel):
    _name = "step.fruit.reception.import"
    _description = "Importar tarjas de recepción de fruta"

    picking_id = fields.Many2one("stock.picking", required=True, readonly=True)
    file = fields.Binary(string="Archivo Excel", required=True)
    filename = fields.Char(string="Nombre de archivo")

    def action_import(self):
        self.ensure_one()
        picking = self.picking_id
        if picking.state == "done" or not picking.step_fruit_reception_kind:
            raise UserError(_("La recepción debe estar abierta y tener tipo de fruta."))
        if not picking.fruit_fundo_id.partner_id:
            raise UserError(_("Seleccione un fundo con productor antes de importar."))
        if not (self.filename or "").lower().endswith(".xlsx"):
            raise UserError(_("Cargue un archivo .xlsx."))
        try:
            workbook = load_workbook(BytesIO(base64.b64decode(self.file)), read_only=True, data_only=True)
            sheet = workbook.active
            rows = sheet.iter_rows(values_only=True)
            headers = tuple(str(value or "").strip() for value in next(rows))
        except Exception as exc:
            raise UserError(_("No fue posible leer el archivo Excel: %s") % exc) from exc
        if headers[:len(_HEADERS)] != _HEADERS:
            raise UserError(_("Las columnas deben ser: %s") % ", ".join(_HEADERS))

        prepared = []
        seen = set(picking.fruit_tag_line_ids.mapped("tag_number"))
        for row_number, row in enumerate(rows, start=2):
            if not any(value is not None for value in row):
                continue
            number, code, qty, kilos, boxes, gross, tare = (list(row) + [None] * 7)[:7]
            number = str(number or "").strip()
            code = str(code or "").strip()
            if not number or number in seen:
                raise UserError(_("Tarja vacía o repetida en fila %s.") % row_number)
            seen.add(number)
            product = self.env["product.product"].search([("default_code", "=", code)], limit=1)
            if not product:
                raise UserError(_("Producto %s de la fila %s no encontrado.") % (code, row_number))
            try:
                qty, kilos, boxes, gross, tare = (float(value or 0) for value in (qty, kilos, boxes, gross, tare))
            except (TypeError, ValueError) as exc:
                raise UserError(_("Cantidad o peso inválido en fila %s.") % row_number) from exc
            if qty <= 0 or kilos <= 0 or boxes < 0 or gross < 0 or tare < 0:
                raise UserError(_("Cantidad y kilos deben ser positivos en fila %s.") % row_number)
            if tare > gross or (gross and round(gross - tare - kilos, 2) != 0):
                raise UserError(_("El peso bruto menos destare no coincide con los kilos en fila %s.") % row_number)
            if self.env["stock.quant.package"].search_count([("name", "=", number)]):
                raise UserError(_("La tarja %s ya existe en Inventario.") % number)
            prepared.append((number, product, qty, kilos, boxes, gross, tare))
        if not prepared:
            raise UserError(_("El archivo no contiene tarjas."))

        kind = "C" if picking.step_fruit_reception_kind == "process" else "E"
        for number, product, qty, kilos, boxes, gross, tare in prepared:
            package = self.env["stock.quant.package"].create({
                "name": number, "is_fruit_tag": True, "step_tag_kind": kind,
                "fundo_id": picking.fruit_fundo_id.id,
                "step_producer_id": picking.fruit_fundo_id.partner_id.id,
                "box_count": boxes,
            })
            self.env["step.packing.picking.tag.line"].create({
                "picking_id": picking.id, "tag_number": number,
                "package_id": package.id, "product_id": product.id,
                "quantity": qty, "kilos": kilos, "box_count": boxes,
                "gross_kg": gross, "tare_kg": tare,
                "uom_id": product.uom_id.id,
            })
        return {"type": "ir.actions.act_window_close"}
